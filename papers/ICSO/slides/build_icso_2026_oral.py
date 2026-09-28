"""Build the ICSO 2026 Paper 293 oral presentation.

The deck reuses the ISSL slide master from the July optical-communication
research-group deck, but all visible slide content is rebuilt from the final
ICSO paper and the 2026-09-12 oral outline.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
SLIDES_DIR = Path(__file__).resolve().parent
WORK_DIR = SLIDES_DIR / "_work" / "assets"
TEMPLATE = ROOT / "papers" / "seminar" / "20260721_optcommrg_takamoto_v3_issl.pptx"
FIG = ROOT / "papers" / "icso" / "figure"
OUT = SLIDES_DIR / "ICSO2026_Paper293_Takamoto_oral_draft.pptx"


NAVY = "082B52"
BLUE = "1B74AA"
MID_BLUE = "4BA3D3"
PALE_BLUE = "EAF5FB"
PALEST_BLUE = "F4F9FC"
ORANGE = "EC9427"
PALE_ORANGE = "FFF4E8"
GREEN = "239571"
PALE_GREEN = "EAF7F2"
RED = "D03B46"
PALE_RED = "FCECEF"
DARK = "232B34"
GRAY = "666F79"
MID_GRAY = "B9C3CD"
LIGHT_GRAY = "F1F4F6"
WHITE = "FFFFFF"

FONT = "Yu Gothic"
MATH_FONT = "Cambria Math"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def remove_all_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def style_text_frame(
    text_frame,
    *,
    size: float,
    color: str = DARK,
    bold: bool = False,
    font: str = FONT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    margin: float = 0.10,
    line_spacing: float = 1.0,
) -> None:
    text_frame.word_wrap = True
    text_frame.vertical_anchor = valign
    text_frame.margin_left = Inches(margin)
    text_frame.margin_right = Inches(margin)
    text_frame.margin_top = Inches(margin)
    text_frame.margin_bottom = Inches(margin)
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 24,
    color: str = DARK,
    bold: bool = False,
    font: str = FONT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    margin: float = 0.10,
    line_spacing: float = 1.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.text = text
    style_text_frame(
        shape.text_frame,
        size=size,
        color=color,
        bold=bold,
        font=font,
        align=align,
        valign=valign,
        margin=margin,
        line_spacing=line_spacing,
    )
    return shape


def add_rich_text(
    slide,
    segments: list[tuple[str, float, str, bool, int]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    align: PP_ALIGN = PP_ALIGN.CENTER,
    valign: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    font: str = MATH_FONT,
):
    """Add one line of rich text; baseline is in DrawingML 1/1000 percent."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = False
    text_frame.vertical_anchor = valign
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    for text, size, color, bold, baseline in segments:
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if baseline:
            run._r.get_or_add_rPr().set("baseline", str(baseline))
    return shape


def add_round_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str = MID_BLUE,
    line_width: float = 1.2,
    radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(radius_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(line_width)
    return shape


def add_card(
    slide,
    title: str,
    body: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    accent: str = BLUE,
    fill: str = PALEST_BLUE,
    title_size: float = 24,
    body_size: float = 20,
    centered: bool = False,
):
    add_round_rect(slide, x, y, w, h, fill=fill, line=accent, line_width=1.1)
    align = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
    add_text(
        slide,
        title,
        x + 0.28,
        y + 0.18,
        w - 0.56,
        0.68,
        size=title_size,
        color=accent,
        bold=True,
        align=align,
        margin=0,
    )
    add_text(
        slide,
        body,
        x + 0.30,
        y + 0.90,
        w - 0.60,
        h - 1.08,
        size=body_size,
        color=DARK,
        align=align,
        valign=MSO_ANCHOR.TOP,
        margin=0,
        line_spacing=1.0,
    )


def add_metric_card(
    slide,
    value: str,
    label: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    accent: str = BLUE,
    fill: str = WHITE,
    value_size: float = 34,
):
    add_round_rect(slide, x, y, w, h, fill=fill, line=accent, line_width=1.2)
    add_text(
        slide,
        value,
        x + 0.18,
        y + 0.18,
        w - 0.36,
        h * 0.50,
        size=value_size,
        color=accent,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_text(
        slide,
        label,
        x + 0.18,
        y + h * 0.58,
        w - 0.36,
        h * 0.25,
        size=17,
        color=GRAY,
        align=PP_ALIGN.CENTER,
        margin=0,
    )


def add_chevron(slide, x: float, y: float, w: float, h: float, color: str = MID_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def add_source(slide, text: str) -> None:
    add_text(
        slide,
        text,
        1.20,
        13.18,
        23.5,
        0.34,
        size=11,
        color=GRAY,
        margin=0,
        valign=MSO_ANCHOR.TOP,
    )


def add_image_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        source_ratio = image.width / image.height
    box_ratio = w / h
    if source_ratio > box_ratio:
        draw_w = w
        draw_h = w / source_ratio
    else:
        draw_h = h
        draw_w = h * source_ratio
    draw_x = x + (w - draw_w) / 2
    draw_y = y + (h - draw_h) / 2
    return slide.shapes.add_picture(
        str(path), Inches(draw_x), Inches(draw_y), Inches(draw_w), Inches(draw_h)
    )


def add_image_exact(slide, path: Path, x: float, y: float, w: float, h: float):
    """Fill a known chart box exactly; used only for deliberately wide plots."""
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def crop_image(path: Path, out_name: str, box: tuple[float, float, float, float]) -> Path:
    """Crop using a normalized (left, top, right, bottom) box."""
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    out = WORK_DIR / out_name
    with Image.open(path) as image:
        left, top, right, bottom = box
        px_box = (
            round(image.width * left),
            round(image.height * top),
            round(image.width * right),
            round(image.height * bottom),
        )
        image.crop(px_box).save(out)
    return out


def make_content_slide(
    prs: Presentation,
    title: str,
    section: str,
    number: str,
    *,
    title_size: float = 32,
):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    for placeholder in list(slide.placeholders):
        remove_shape(placeholder)
    add_text(slide, section, 0.60, 0.22, 17.5, 0.72, size=20, color=GRAY, margin=0)
    add_text(
        slide,
        title,
        1.70,
        1.42,
        24.15,
        1.70,
        size=title_size,
        color=NAVY,
        bold=True,
        margin=0,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        str(number),
        24.95,
        13.70,
        1.10,
        0.55,
        size=18,
        color=GRAY,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )
    return slide


def add_table(
    slide,
    rows: list[list[str]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    col_widths: list[float] | None = None,
    font_size: float = 17,
    header_size: float = 18,
    highlight_last: bool = False,
):
    table = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    if col_widths:
        total = sum(col_widths)
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = Inches(w * width / total)
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            if row_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(NAVY)
                color = WHITE
                bold = True
                size = header_size
            elif highlight_last and row_idx == len(rows) - 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(PALE_GREEN)
                color = NAVY
                bold = True
                size = font_size
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(WHITE if row_idx % 2 else LIGHT_GRAY)
                color = DARK
                bold = False
                size = font_size
            style_text_frame(
                cell.text_frame,
                size=size,
                color=color,
                bold=bold,
                align=PP_ALIGN.LEFT if col_idx == 0 else PP_ALIGN.CENTER,
                margin=0.08,
            )
    return table


def add_comparison_bar(
    slide,
    label: str,
    value: float,
    value_label: str,
    max_value: float,
    x: float,
    y: float,
    w: float,
    *,
    color: str,
    min_fraction: float = 0.015,
):
    add_text(slide, label, x, y, 4.4, 0.70, size=22, color=DARK, margin=0)
    track_x = x + 4.6
    add_round_rect(
        slide,
        track_x,
        y + 0.02,
        w - 6.3,
        0.66,
        fill=LIGHT_GRAY,
        line=LIGHT_GRAY,
        line_width=0,
        radius_shape=MSO_SHAPE.RECTANGLE,
    )
    bar_w = max((w - 6.3) * value / max_value, (w - 6.3) * min_fraction)
    add_round_rect(
        slide,
        track_x,
        y + 0.02,
        bar_w,
        0.66,
        fill=color,
        line=color,
        line_width=0,
        radius_shape=MSO_SHAPE.RECTANGLE,
    )
    add_text(
        slide,
        value_label,
        x + w - 1.55,
        y,
        1.55,
        0.70,
        size=22,
        color=color,
        bold=True,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )


def build_deck() -> Path:
    prs = Presentation(TEMPLATE)
    remove_all_slides(prs)
    prs.core_properties.title = (
        "Hierarchical Prediction and Feedforward Correction of Time-Varying Thermal LOS Bias"
    )
    prs.core_properties.subject = "ICSO 2026 oral presentation, Paper 293"
    prs.core_properties.author = "Hideki Takamoto"
    prs.core_properties.comments = "Generated from the final ICSO paper and oral-slide outline."

    # Slide 1 — title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for placeholder in list(slide.placeholders):
        remove_shape(placeholder)
    add_text(
        slide,
        "Hierarchical Prediction and Feedforward Correction of\n"
        "Time-Varying Thermal Line-of-Sight Bias for Coarse Acquisition\n"
        "in Satellite Optical Communications",
        10.60,
        7.65,
        14.65,
        2.70,
        size=29,
        color=NAVY,
        bold=True,
        margin=0,
        valign=MSO_ANCHOR.TOP,
        line_spacing=0.90,
    )
    add_text(
        slide,
        "Hideki Takamoto, Kazuki Takashima, Yuki Kusano, Satoshi Ikari, and Ryu Funase",
        10.60,
        11.43,
        14.65,
        0.80,
        size=23,
        color=DARK,
        bold=True,
        margin=0,
    )
    add_text(
        slide,
        "Department of Aeronautics and Astronautics, The University of Tokyo",
        10.60,
        12.25,
        14.65,
        0.60,
        size=18,
        color=GRAY,
        margin=0,
    )
    add_text(slide, "ICSO 2026", 18.05, 5.90, 6.15, 0.50, size=22, color=GRAY, margin=0)
    add_text(slide, "Paper 293 · Oral", 18.05, 6.43, 6.15, 0.50, size=22, color=GRAY, margin=0)
    add_text(
        slide,
        "PAT and Receiver Technologies",
        18.05,
        6.96,
        6.15,
        0.48,
        size=18,
        color=GRAY,
        margin=0,
    )
    add_text(slide, "1", 24.95, 13.70, 1.10, 0.55, size=18, color=GRAY, align=PP_ALIGN.RIGHT, margin=0)

    # Slide 2 — problem and positioning
    slide = make_content_slide(
        prs,
        "Thermal LOS bias can dominate the coarse-acquisition scan before optical feedback is available",
        "Background and positioning",
        "2",
        title_size=29,
    )
    cards = [
        (
            "Before optical feedback",
            "Attitude, orbit, alignment, and thermal errors define the search region.",
            BLUE,
            PALE_BLUE,
        ),
        (
            "Thermoelastic bias",
            "Known illumination and dissipation drive bus deformation.\n\n150 µrad to >1 mrad",
            ORANGE,
            PALE_ORANGE,
        ),
        (
            "Feedforward scan center",
            "Predict the thermal component before acquisition and subtract it from the scan center.",
            GREEN,
            PALE_GREEN,
        ),
    ]
    for idx, (title, body, accent, fill) in enumerate(cards):
        x = 1.20 + idx * 8.20
        add_card(
            slide,
            title,
            body,
            x,
            3.35,
            7.05,
            3.20,
            accent=accent,
            fill=fill,
            title_size=24,
            body_size=20,
            centered=True,
        )
        if idx < 2:
            add_chevron(slide, x + 7.32, 4.45, 0.55, 0.75, MID_BLUE)
    add_text(slide, "Representative prior work", 1.30, 6.85, 8.0, 0.48, size=17, color=GRAY, margin=0)
    position_rows = [
        ["Research stream", "Thermal-state LOS prediction", "Acquisition evaluation"],
        ["Optical communication / PAT", "Limited", "Yes"],
        ["Earth-observation & deep-space instruments", "Yes", "No"],
        ["This work", "Yes", "Yes"],
    ]
    add_table(
        slide,
        position_rows,
        1.20,
        7.35,
        24.25,
        4.70,
        col_widths=[2.2, 1.25, 1.25],
        font_size=18,
        header_size=17,
        highlight_last=True,
    )
    add_text(
        slide,
        "This work connects thermal-state prediction to acquisition performance.",
        4.25,
        12.12,
        18.3,
        0.66,
        size=23,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_source(
        slide,
        "Representative sources: Kaushal 2017; Shi 2023; Zhang 2025; Hu 2022; Li 2025; Turella 2019/2021.",
    )

    # Slide 3 — LOS definition
    slide = make_content_slide(
        prs,
        "The relevant far-field quantity is the LCT optical-axis rotation relative to the STT",
        "Problem definition",
        "3",
        title_size=30,
    )
    add_image_contain(slide, FIG / "fig_los_definition.png", 1.15, 3.25, 15.2, 9.65)
    add_round_rect(slide, 17.15, 3.78, 8.15, 2.10, fill=PALE_BLUE, line=BLUE, line_width=1.3)
    add_rich_text(
        slide,
        [
            ("θ", 38, NAVY, True, 0),
            ("th", 24, NAVY, True, -25000),
            (" = θ", 38, NAVY, True, 0),
            ("LCT", 24, NAVY, True, -25000),
            (" − θ", 38, NAVY, True, 0),
            ("STT", 24, NAVY, True, -25000),
        ],
        17.45,
        4.15,
        7.55,
        1.25,
    )
    add_card(
        slide,
        "Attitude reference",
        "Thermal rotation of the STT is absorbed into the estimated attitude frame.",
        17.15,
        6.35,
        8.15,
        2.35,
        accent=BLUE,
        fill=WHITE,
        title_size=22,
        body_size=19,
    )
    add_card(
        slide,
        "Communication axis",
        "The LCT rotation relative to that frame appears directly as scan-center error.",
        17.15,
        9.03,
        8.15,
        2.35,
        accent=GREEN,
        fill=WHITE,
        title_size=22,
        body_size=19,
    )
    add_text(
        slide,
        "Centerline translation is negligible for the far-field link.",
        17.35,
        11.73,
        7.75,
        0.60,
        size=17,
        color=GRAY,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_source(slide, "Paper §3, Figure 2.")

    # Slide 4 — analysis setup
    slide = make_content_slide(
        prs,
        "A box-structure LEO bus is analysed across 21 operating conditions",
        "Thermoelastic analysis",
        "4",
        title_size=31,
    )
    add_round_rect(slide, 1.10, 3.30, 16.40, 9.25, fill=WHITE, line=MID_GRAY, line_width=0.8)
    add_image_contain(slide, FIG / "fig_td_tdall_femap.png", 1.30, 3.52, 16.00, 8.80)
    add_card(
        slide,
        "Spacecraft model",
        "0.6 × 0.6 × 1.0 m\nA5052, reference 24 °C\nSTT on PZ · LCT on MZ",
        18.05,
        3.35,
        7.20,
        2.55,
        accent=BLUE,
        fill=PALEST_BLUE,
        title_size=22,
        body_size=18,
    )
    add_card(
        slide,
        "21 conditions",
        "Sun face: MX / MY / PX / PY\nDissipation · surface · orbit variants",
        18.05,
        6.20,
        7.20,
        2.25,
        accent=ORANGE,
        fill=PALE_ORANGE,
        title_size=22,
        body_size=18,
    )
    add_card(
        slide,
        "Common analysis chain",
        "Thermal Desktop temperature field\n→ Femap rotations\n→ STT–LCT relative LOS",
        18.05,
        8.75,
        7.20,
        2.65,
        accent=GREEN,
        fill=PALE_GREEN,
        title_size=22,
        body_size=18,
    )
    add_text(
        slide,
        "One fixed structure; operating conditions vary across cases.",
        18.20,
        11.65,
        6.90,
        0.75,
        size=18,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_source(slide, "Paper §4–5, Tables 2–4, Figure 3.")

    # Slide 5 — observations
    temp_my = crop_image(
        FIG / "fig_temp_field_case04.png",
        "case04_panel_my.png",
        (0.03, 0.31, 0.99, 0.67),
    )
    los_y = crop_image(
        FIG / "p1_far_field_los_case04.png",
        "case04_los_y.png",
        (0.03, 0.47, 0.99, 0.995),
    )
    slide = make_content_slide(
        prs,
        "Thermal LOS follows the orbit; its dominant axis and DC offset follow the sun face",
        "Observed structure",
        "5",
        title_size=30,
    )
    add_text(slide, "MY-panel temperature · Case 04", 1.25, 3.23, 11.60, 0.55, size=19, color=BLUE, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "STT-relative LOS, dominant y axis", 13.65, 3.23, 11.60, 0.55, size=19, color=GREEN, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_image_exact(slide, temp_my, 1.15, 4.35, 11.75, 4.80)
    add_image_exact(slide, los_y, 13.55, 4.35, 11.75, 4.80)
    add_card(
        slide,
        "Same orbital period",
        "Temperature and thermal LOS repeat together.",
        1.20,
        10.35,
        7.55,
        2.10,
        accent=BLUE,
        fill=PALEST_BLUE,
        title_size=21,
        body_size=17,
        centered=True,
    )
    add_card(
        slide,
        "Sun face sets the axis",
        "MY/PY → y dominant\nMX/PX → x dominant",
        9.55,
        10.35,
        7.55,
        2.10,
        accent=ORANGE,
        fill=PALE_ORANGE,
        title_size=21,
        body_size=17,
        centered=True,
    )
    add_card(
        slide,
        "Raw RMS: 150–1280 µrad",
        "Dissipation mainly shifts the case DC bias.",
        17.90,
        10.35,
        7.55,
        2.10,
        accent=GREEN,
        fill=PALE_GREEN,
        title_size=21,
        body_size=17,
        centered=True,
    )
    add_source(slide, "Paper §5, Figure 4. Plots are cropped from the final-paper source figures.")

    # Slide 6 — model
    slide = make_content_slide(
        prs,
        "Orbital variation is a(sun)·ΔT; the DC term is predicted from sun face and dissipation",
        "Hierarchical sun-face ΔT model",
        "6",
        title_size=29,
    )
    add_card(
        slide,
        "Onboard inputs",
        "Panel-center temperature difference\nΔT(t) = T(sun face) − T(opposite)\n\nSun face + dissipation flags",
        1.20,
        3.60,
        6.20,
        3.85,
        accent=BLUE,
        fill=PALE_BLUE,
        title_size=23,
        body_size=19,
        centered=True,
    )
    add_chevron(slide, 7.70, 4.95, 0.72, 1.05, MID_BLUE)
    add_round_rect(slide, 8.75, 3.47, 8.95, 4.12, fill=WHITE, line=NAVY, line_width=1.5)
    add_text(slide, "Hierarchical prediction", 9.10, 3.78, 8.25, 0.62, size=23, color=NAVY, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_rich_text(
        slide,
        [
            ("θ", 32, NAVY, True, 0),
            ("dom", 20, NAVY, True, -25000),
            ("(t) ≈ b", 32, NAVY, True, 0),
            ("case", 20, NAVY, True, -25000),
            (" + a(sun)·ΔT(t)", 32, NAVY, True, 0),
        ],
        9.10,
        4.57,
        8.25,
        1.02,
    )
    add_rich_text(
        slide,
        [
            ("b", 24, ORANGE, True, 0),
            ("case", 16, ORANGE, True, -25000),
            (" ≈ b", 24, ORANGE, True, 0),
            ("0", 16, ORANGE, True, -25000),
            ("(sun) + c", 24, ORANGE, True, 0),
            ("PROP", 15, ORANGE, True, -25000),
            (" I", 24, ORANGE, True, 0),
            ("PROP", 15, ORANGE, True, -25000),
            (" + c", 24, ORANGE, True, 0),
            ("PCDU", 15, ORANGE, True, -25000),
            (" I", 24, ORANGE, True, 0),
            ("PCDU", 15, ORANGE, True, -25000),
        ],
        9.10,
        5.83,
        8.25,
        0.78,
    )
    add_text(slide, "16 fixed coefficients", 11.30, 6.73, 3.85, 0.45, size=17, color=GRAY, align=PP_ALIGN.CENTER, margin=0)
    add_chevron(slide, 18.00, 4.95, 0.72, 1.05, MID_BLUE)
    add_card(
        slide,
        "Predicted thermal LOS",
        "Two-axis scan-center correction\ncomputed before acquisition",
        19.05,
        3.60,
        6.20,
        3.85,
        accent=GREEN,
        fill=PALE_GREEN,
        title_size=23,
        body_size=20,
        centered=True,
    )
    add_card(
        slide,
        "Within-orbit variation",
        "A shared sensitivity a(sun) maps ΔT(t) to the dominant-axis time variation.",
        2.05,
        8.35,
        10.75,
        2.55,
        accent=BLUE,
        fill=WHITE,
        title_size=23,
        body_size=19,
    )
    add_card(
        slide,
        "Across-condition DC",
        "Sun face and operating flags predict the case DC bias across conditions.",
        13.85,
        8.35,
        10.75,
        2.55,
        accent=ORANGE,
        fill=WHITE,
        title_size=23,
        body_size=19,
    )
    add_text(
        slide,
        "The first-order ΔT–LOS relation is known; the contribution here is bus-relative STT–LCT LOS, cross-condition DC prediction, and acquisition use.",
        2.10,
        11.45,
        22.45,
        0.95,
        size=19,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_source(slide, "Paper §6, Eqs. 3–6, Figure 5.")

    # Slide 7 — cross-case accuracy
    case08_plot = crop_image(
        FIG / "p2_bcase_true_vs_pred_case08.png",
        "case08_prediction_without_log_title.png",
        (0.00, 0.055, 1.00, 0.99),
    )
    slide = make_content_slide(
        prs,
        "Nested leave-one-case-out RMSE is 4.9 µrad against 615 µrad raw RMS",
        "Cross-condition validation",
        "7",
        title_size=31,
    )
    add_image_contain(slide, FIG / "p3_a_emp_by_sunface.png", 1.10, 3.25, 11.75, 7.15)
    add_image_contain(slide, FIG / "p3_b_emp_vs_b_pred.png", 13.05, 3.25, 12.40, 7.15)
    add_metric_card(slide, "28–31 µrad/K", "shared sensitivity magnitude", 1.30, 10.62, 7.25, 2.00, accent=BLUE, fill=PALEST_BLUE, value_size=30)
    add_metric_card(slide, "3.8 µrad", "LOO case-bias RMSE", 9.70, 10.62, 7.25, 2.00, accent=ORANGE, fill=PALE_ORANGE, value_size=32)
    add_metric_card(slide, "4.9 vs 615 µrad", "median test RMSE vs raw RMS", 18.10, 10.62, 7.25, 2.00, accent=GREEN, fill=PALE_GREEN, value_size=30)
    add_source(slide, "21 cases; the test case is excluded from shared sensitivity and both Level-2 axis-bias estimates. Paper §6, Figure 6.")

    # Slide 8 — example prediction
    slide = make_content_slide(
        prs,
        "Case 08 (PY, all units on): 1250 µrad raw RMS falls to 3.9 µrad test RMSE",
        "Time-series example",
        "8",
        title_size=30,
    )
    add_round_rect(slide, 1.00, 3.20, 14.20, 9.55, fill=WHITE, line=MID_GRAY, line_width=0.7)
    add_image_contain(slide, case08_plot, 1.20, 3.35, 13.80, 9.20)
    add_metric_card(slide, "1250 µrad", "raw dominant-axis RMS", 16.05, 3.75, 8.90, 2.10, accent=RED, fill=PALE_RED, value_size=34)
    add_metric_card(slide, "3.9 µrad", "nested-LOO test RMSE", 16.05, 6.25, 8.90, 2.10, accent=GREEN, fill=PALE_GREEN, value_size=34)
    add_card(
        slide,
        "What the model captures",
        "The predicted case DC sets the offset; shared a(sun)·ΔT(t) follows the repeating orbital variation.",
        16.05,
        8.75,
        8.90,
        2.70,
        accent=BLUE,
        fill=WHITE,
        title_size=22,
        body_size=19,
    )
    add_text(slide, "Prediction tracks the largest thermal case with a two-order-of-magnitude reduction.", 16.30, 11.83, 8.40, 0.62, size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "Paper §6, Figure 7.")

    # Slide 9 — PAT connection
    slide = make_content_slide(
        prs,
        "The prediction is subtracted from the scan center before the rectangular spiral starts",
        "Coarse-acquisition evaluation",
        "9",
        title_size=30,
    )
    flow = [
        ("TD / Femap truth", "Thermal LOS\ntime series", BLUE, PALE_BLUE),
        ("Hierarchical model", "Predicted thermal LOS\nfrom ΔT and operating flags", ORANGE, PALE_ORANGE),
        ("Scan-center correction", "Corrected center\n= nominal − predicted thermal LOS", GREEN, PALE_GREEN),
        ("Rectangular spiral", "First point inside\ndetection radius", NAVY, PALEST_BLUE),
    ]
    for idx, (title, body, accent, fill) in enumerate(flow):
        x = 0.95 + idx * 6.50
        add_card(slide, title, body, x, 3.55, 5.55, 3.25, accent=accent, fill=fill, title_size=21, body_size=18, centered=True)
        if idx < 3:
            add_chevron(slide, x + 5.77, 4.68, 0.46, 0.72, MID_BLUE)
    add_text(
        slide,
        "Acquisition time is used as a practical proxy for residual initial pointing uncertainty.",
        4.25,
        7.25,
        18.15,
        0.70,
        size=23,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    params = [
        ("±1600 µrad", "scan range"),
        ("120 µrad", "scan step"),
        ("150 µrad", "detection radius"),
        ("0.1 s", "dwell / point"),
        ("27 × 27", "maximum grid"),
    ]
    for idx, (value, label) in enumerate(params):
        add_metric_card(slide, value, label, 1.15 + idx * 5.00, 8.45, 4.35, 2.25, accent=BLUE if idx < 2 else GREEN, fill=WHITE, value_size=27)
    add_text(slide, "No coverage holes: 120/√2 ≈ 85 µrad < 150 µrad.", 6.15, 11.18, 14.40, 0.55, size=18, color=GRAY, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "Slew, settling, and stochastic detection are not modeled.", 6.15, 11.80, 14.40, 0.55, size=18, color=RED, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "Paper §7.1, Table 6, Figures 8–9; beacon field basis: Shi et al. 2023.")

    # Slide 10 — thermal-only result
    slide = make_content_slide(
        prs,
        "With thermal error only, mean acquisition time falls from 14.6 s to 0.10 s",
        "PAT results · thermal only",
        "10",
        title_size=31,
    )
    add_text(slide, "Mean acquisition time", 1.35, 3.45, 14.0, 0.65, size=23, color=NAVY, bold=True, margin=0)
    add_comparison_bar(slide, "No correction", 14.6, "14.6 s", 14.6, 1.35, 4.35, 16.90, color=RED)
    add_comparison_bar(slide, "Hierarchical ΔT", 0.10, "0.10 s", 14.6, 1.35, 5.65, 16.90, color=GREEN)
    add_comparison_bar(slide, "Thermal truth", 0.10, "0.10 s", 14.6, 1.35, 6.95, 16.90, color=BLUE)
    add_metric_card(slide, "9.3 µrad", "mean two-axis thermal residual", 19.10, 3.85, 6.15, 2.25, accent=GREEN, fill=PALE_GREEN, value_size=33)
    add_metric_card(slide, "100%", "success for all three methods", 19.10, 6.55, 6.15, 2.25, accent=BLUE, fill=PALE_BLUE, value_size=35)
    add_card(
        slide,
        "Why one scan point is enough",
        "9.3 µrad residual ≪ 150 µrad detection radius.\nThe prediction reaches the thermal-truth upper bound.",
        2.00,
        9.15,
        22.70,
        2.60,
        accent=GREEN,
        fill=WHITE,
        title_size=23,
        body_size=21,
        centered=True,
    )
    add_text(slide, "14 COLD baseline-surface cases · 301 epochs × 3 orbits per case", 5.00, 12.15, 16.70, 0.52, size=18, color=GRAY, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "Paper §7.3, Table 7.")

    # Slide 11 — nonthermal result
    slide = make_content_slide(
        prs,
        "With nonthermal error, the remaining search is no longer thermal: 19.2 s to 5.45 s",
        "PAT results · with nonthermal error",
        "11",
        title_size=30,
    )
    add_text(slide, "Mean acquisition time", 1.35, 3.45, 14.0, 0.65, size=23, color=NAVY, bold=True, margin=0)
    add_comparison_bar(slide, "No correction", 19.2, "19.2 s", 19.2, 1.35, 4.35, 16.90, color=RED)
    add_comparison_bar(slide, "Hierarchical ΔT", 5.45, "5.45 s", 19.2, 1.35, 5.75, 16.90, color=GREEN)
    add_metric_card(slide, "−72%", "mean acquisition time", 19.10, 3.85, 6.15, 2.25, accent=GREEN, fill=PALE_GREEN, value_size=36)
    add_metric_card(slide, "98 → 100%", "acquisition success", 19.10, 6.55, 6.15, 2.25, accent=BLUE, fill=PALE_BLUE, value_size=31)
    add_card(
        slide,
        "After correction",
        "Mean initial error = 448 µrad\n≈ synthesized nonthermal error floor",
        1.50,
        8.30,
        11.00,
        2.65,
        accent=ORANGE,
        fill=PALE_ORANGE,
        title_size=23,
        body_size=21,
        centered=True,
    )
    add_card(
        slide,
        "Largest benefit: PY sun face",
        "37.3–39.5 s and 88–97%\n→ 1.3–1.8 s and 100%",
        14.15,
        8.30,
        11.00,
        2.65,
        accent=GREEN,
        fill=PALE_GREEN,
        title_size=23,
        body_size=21,
        centered=True,
    )
    add_text(
        slide,
        "The method removes predictable thermal bias; it does not claim to eliminate total pointing error.",
        3.10,
        11.48,
        20.45,
        0.75,
        size=22,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_source(slide, "One nonthermal-error seed per case; not a Monte Carlo ensemble. Paper §7.2–7.3, Table 7.")

    # Slide 12 — limitations
    slide = make_content_slide(
        prs,
        "The present evaluation is numerical and structure-specific; several floors remain",
        "Scope and limitations",
        "12",
        title_size=31,
    )
    limitations = [
        (
            "Same simulation chain",
            "Truth and reduced model use the same TD/Femap mechanics. LOO removes cases, not the shared structure.",
            BLUE,
            PALE_BLUE,
        ),
        (
            "LOS extraction proxy",
            "Representative node rotations are used; a rigid-body fit over mounting interfaces remains to be tested.",
            ORANGE,
            PALE_ORANGE,
        ),
        (
            "Unmodeled operating factors",
            "Black coating, HOT orbit, and half-power dissipation raise the residual floor and are not explicit in Level 2.",
            RED,
            PALE_RED,
        ),
        (
            "Simplified acquisition",
            "Nonthermal error uses one seed; scan dynamics and stochastic detection are neglected; no ground test yet.",
            GREEN,
            PALE_GREEN,
        ),
    ]
    for idx, (title, body, accent, fill) in enumerate(limitations):
        col = idx % 2
        row = idx // 2
        add_card(
            slide,
            title,
            body,
            1.30 + col * 12.45,
            3.45 + row * 4.05,
            11.60,
            3.45,
            accent=accent,
            fill=fill,
            title_size=24,
            body_size=19,
        )
    add_round_rect(slide, 3.00, 11.68, 20.65, 0.95, fill=WHITE, line=NAVY, line_width=1.0)
    add_text(
        slide,
        "The measured sensitivity (~30 µrad/K) is specific to this structure, placement, and LOS definition; another spacecraft requires re-identification.",
        3.35,
        11.86,
        19.95,
        0.55,
        size=19,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_source(slide, "Paper §6 limitations, §7.4, §8.")

    # Slide 13 — conclusion
    slide = make_content_slide(
        prs,
        "Predictable thermal bias can be removed before acquisition; nonthermal error sets the remaining search",
        "Conclusion",
        "13",
        title_size=30,
    )
    add_metric_card(slide, "4.9 vs 615 µrad", "nested-LOO RMSE vs raw RMS", 1.20, 3.75, 7.55, 3.00, accent=BLUE, fill=PALE_BLUE, value_size=31)
    add_metric_card(slide, "14.6 → 0.10 s", "thermal error only", 9.55, 3.75, 7.55, 3.00, accent=GREEN, fill=PALE_GREEN, value_size=33)
    add_metric_card(slide, "19.2 → 5.45 s", "with nonthermal error", 17.90, 3.75, 7.55, 3.00, accent=ORANGE, fill=PALE_ORANGE, value_size=33)
    add_round_rect(slide, 1.45, 7.55, 23.75, 3.30, fill=WHITE, line=NAVY, line_width=1.5)
    add_text(
        slide,
        "Bus-relative STT–LCT thermal LOS",
        2.00,
        8.00,
        7.00,
        0.70,
        size=24,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_chevron(slide, 9.28, 8.06, 0.65, 0.62, MID_BLUE)
    add_text(
        slide,
        "Shared coefficients across conditions",
        10.10,
        8.00,
        7.00,
        0.70,
        size=24,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_chevron(slide, 17.38, 8.06, 0.65, 0.62, MID_BLUE)
    add_text(
        slide,
        "Feedforward coarse acquisition",
        18.15,
        8.00,
        6.45,
        0.70,
        size=24,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_text(
        slide,
        "Under the evaluated conditions, this reduces search burden before optical feedback is available.",
        3.40,
        9.35,
        19.85,
        0.85,
        size=25,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_text(slide, "Thank you", 9.00, 11.55, 8.70, 0.75, size=28, color=GRAY, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "ICSO 2026 · Paper 293")

    # Backup B1 — case matrix
    slide = make_content_slide(prs, "The 21-case matrix separates sun face, dissipation, surface, and orbit effects", "BACKUP · Case matrix", "B1", title_size=30)
    case_rows = [
        ["Case(s)", "Purpose", "Conditions"],
        ["04–06, 08–09", "Sun face / baseline", "MX/MY/PX/PY; all units or STT/LCT only"],
        ["10", "Thermal environment", "MY; HOT orbit with continuous illumination"],
        ["11–12", "Surface properties", "MY; Black sun face or Alodine surfaces"],
        ["13–21", "Dissipation modes", "PROP only / PCDU only / no additional dissipation"],
        ["22", "Continuous power", "MY; half PROP power (12.5 W) + PCDU"],
        ["23–24", "MX dissipation", "MX; PROP only or PCDU only"],
        ["25", "Orbit condition", "MY; LTAN18, 693 km"],
    ]
    add_table(slide, case_rows, 1.15, 3.35, 24.35, 8.75, col_widths=[0.95, 1.55, 3.0], font_size=17, header_size=19)
    add_text(slide, "Evaluated IDs: 04–06 and 08–25. Cases 01–03 and 07 were MZ/setup runs and are excluded.", 2.35, 12.35, 21.95, 0.55, size=18, color=GRAY, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "Paper Table 4.")

    # Backup B2 — scan geometry
    slide = make_content_slide(prs, "A 120 µrad step with a 150 µrad detection radius leaves no coverage holes", "BACKUP · Scan geometry", "B2", title_size=30)
    add_image_contain(slide, FIG / "fig_rectangular_scan.png", 1.20, 3.10, 13.80, 9.85)
    add_card(slide, "Coverage condition", "Half-diagonal = 120/√2 ≈ 85 µrad\n85 µrad < 150 µrad detection radius", 15.60, 3.60, 9.45, 2.80, accent=GREEN, fill=PALE_GREEN, title_size=24, body_size=22, centered=True)
    add_card(slide, "Full search", "Range: ±1600 µrad\nGrid: 27 × 27 = 729 points\nMaximum: 72.9 s", 15.60, 6.90, 9.45, 2.85, accent=BLUE, fill=PALE_BLUE, title_size=24, body_size=21, centered=True)
    add_card(slide, "Interpretation", "The 150 µrad radius represents a coarse beacon field, not the narrow communication beam.", 15.60, 10.25, 9.45, 2.15, accent=ORANGE, fill=PALE_ORANGE, title_size=22, body_size=18)
    add_source(slide, "Paper §7.1, Table 6, Figure 9; field-of-view basis: Shi et al. 2023.")

    # Backup B3 — nonthermal errors
    slide = make_content_slide(prs, "The nonthermal synthesis combines orbit, alignment, attitude, and drift errors", "BACKUP · Nonthermal error", "B3", title_size=30)
    items = [
        ("Orbit prediction", "Latest Sentinel-1 TLE propagated by SGP4\nvs precise POEORB, projected to link transverse plane", BLUE, PALE_BLUE),
        ("Alignment residual", "Constant two-axis residual\n1σ = 50 µrad", ORANGE, PALE_ORANGE),
        ("Attitude error", "Random two-axis determination/control error\n1σ = 50 µrad", GREEN, PALE_GREEN),
        ("Low-frequency drift", "Amplitude 30 µrad\nPeriod 900 s", RED, PALE_RED),
    ]
    for idx, (title, body, accent, fill) in enumerate(items):
        add_card(slide, title, body, 1.20 + (idx % 2) * 12.35, 3.45 + (idx // 2) * 3.65, 11.55, 3.05, accent=accent, fill=fill, title_size=23, body_size=19, centered=True)
    add_round_rect(slide, 2.45, 11.05, 21.75, 1.35, fill=WHITE, line=NAVY, line_width=1.1)
    add_text(slide, "Thermal LOS and orbit error both contain near-orbital-period content; frequency alone cannot isolate the thermal term after acquisition.", 2.85, 11.30, 20.95, 0.85, size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "One fixed seed per case; not a spacecraft-specific error budget or Monte Carlo ensemble. Paper §7.2.")

    # Backup B4 — residual Fourier
    slide = make_content_slide(prs, "A preliminary residual Fourier update preserves first-orbit FF and lowers later-orbit search", "BACKUP · Preliminary residual update", "B4", title_size=29)
    add_image_contain(slide, FIG / "fig_residual_update_comparison.png", 1.05, 3.20, 16.20, 8.55)
    add_card(slide, "Causal update", "K = 2 Fourier coefficients are fitted in orbit n and applied to orbit n+1.", 17.75, 3.70, 7.40, 2.45, accent=BLUE, fill=PALE_BLUE, title_size=22, body_size=18)
    add_card(slide, "Why FF first", "Direct Fourier is uncorrected in orbit 0: Case 16 requires 39.5 s.", 17.75, 6.55, 7.40, 2.45, accent=RED, fill=PALE_RED, title_size=22, body_size=18)
    add_card(slide, "Preliminary only", "Two cases; dense 60.5 s samples; failed-point residuals and receiver realization are not validated.", 17.75, 9.40, 7.40, 2.55, accent=ORANGE, fill=PALE_ORANGE, title_size=22, body_size=17)
    add_source(slide, "Paper §7.4, Figure 10.")

    # Backup B5 — attitude extraction
    slide = make_content_slide(prs, "Representative node rotations are a proxy for equipment attitude", "BACKUP · LOS extraction", "B5", title_size=31)
    add_image_contain(slide, FIG / "fig_los_definition.png", 1.20, 3.20, 13.80, 9.55)
    add_card(slide, "Implemented", "Thermal rotations at representative STT and LCT center nodes are differenced in the transverse axes.", 15.60, 3.60, 9.30, 2.65, accent=BLUE, fill=PALE_BLUE, title_size=24, body_size=19)
    add_card(slide, "Not yet implemented", "Least-squares rigid-body attitude fitted over each mounting interface.", 15.60, 6.75, 9.30, 2.45, accent=ORANGE, fill=PALE_ORANGE, title_size=24, body_size=20)
    add_card(slide, "Implication", "The present LOS definition and extraction method are shared by truth and reduced model.", 15.60, 9.70, 9.30, 2.25, accent=GREEN, fill=PALE_GREEN, title_size=23, body_size=19)
    add_source(slide, "Paper §4, analysis procedure.")

    # Backup B6 — hard cases
    slide = make_content_slide(prs, "The largest residuals identify missing variables rather than model universality", "BACKUP · Hard cases", "B6", title_size=30)
    hard_cases = [
        ("Black coating", "≈13 µrad test RMSE\nSurface properties change orbital amplitude and residual floor.", RED, PALE_RED),
        ("HOT orbit", "Several-µrad DC residual\nContinuous illumination is not explicit in Level 2.", ORANGE, PALE_ORANGE),
        ("Half PROP power", "≈16 µrad test RMSE\nA binary dissipation flag cannot represent continuous power.", BLUE, PALE_BLUE),
    ]
    for idx, (title, body, accent, fill) in enumerate(hard_cases):
        add_card(slide, title, body, 1.20 + idx * 8.15, 4.05, 7.40, 4.10, accent=accent, fill=fill, title_size=25, body_size=21, centered=True)
    add_round_rect(slide, 2.80, 9.25, 21.05, 2.10, fill=WHITE, line=GREEN, line_width=1.4)
    add_text(slide, "Next Level-2 variables", 3.20, 9.55, 5.70, 0.60, size=23, color=GREEN, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "surface state  ·  orbit/thermal environment  ·  continuous dissipation", 8.75, 9.48, 14.45, 0.78, size=23, color=NAVY, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "The model structure is portable; the present coefficients are not claimed to be universal.", 3.50, 11.75, 19.70, 0.65, size=21, color=GRAY, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "Paper §6 discussion and limitations.")

    # Backup B7 — constant bias only placeholder
    slide = make_content_slide(prs, "Constant-bias-only PAT comparison is reserved for a controlled ablation", "BACKUP · Pending ablation", "B7", title_size=30)
    add_round_rect(slide, 3.10, 4.00, 20.45, 5.70, fill=LIGHT_GRAY, line=MID_GRAY, line_width=1.2)
    add_text(slide, "PENDING ANALYSIS", 5.10, 4.75, 16.45, 0.95, size=36, color=GRAY, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "Set ΔT(t) = 0 while retaining the predicted case DC bias.\nThis will quantify how much acquisition improvement comes from DC correction versus orbital variation.", 5.10, 6.05, 16.45, 1.85, size=24, color=DARK, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "No value from this pending ablation is used in the main-slide conclusions.", 5.10, 8.40, 16.45, 0.60, size=20, color=RED, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "Reserved slot from the oral-slide outline; not yet computed.")

    # Backup B8 — related work
    slide = make_content_slide(prs, "This work connects thermal-state LOS prediction to optical-communication acquisition", "BACKUP · Prior work and positioning", "B8", title_size=29)
    prior_rows = [
        ["Study", "System / domain", "Thermal treatment", "Acquisition"],
        ["Riesing 2023", "CubeSat optical PAT", "Not modeled; absorbed by attitude system", "Yes, in orbit"],
        ["Rüddenklau 2026", "Optical terminal", "FF of attitude / mounting errors", "Yes"],
        ["Shi 2023", "Optical-comm structure", "Reduced by structural design", "Yes"],
        ["Zhang 2025", "Inter-satellite pointing", "Finite-element assessment / guidance", "Impact evaluated"],
        ["Hu 2022 / Li 2025", "GEO / LEO Earth observation", "Periodic model / neural network", "No"],
        ["Turella 2019/2021", "JUICE/JANUS camera", "Wall-to-wall ΔT proportional model", "No"],
        ["This work", "Optical-comm spacecraft bus", "Hierarchical sun-face ΔT prediction", "Yes"],
    ]
    add_table(slide, prior_rows, 0.95, 3.28, 24.75, 8.95, col_widths=[1.25, 1.60, 2.35, 1.10], font_size=14, header_size=17, highlight_last=True)
    add_text(slide, "Novelty is not the first-order temperature–LOS relation itself; it is the bus-relative STT–LCT application, cross-condition DC prediction, and acquisition connection.", 2.30, 12.35, 22.05, 0.62, size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_source(slide, "Paper §2 and related-work overview table.")

    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_deck()
    print(path)

"""Build the 2026-09-28 lab-wide seminar deck (Japanese).

The deck starts from the July optical-communication research-group deck
(`20260721_optcommrg_takamoto_v3_issl.pptx`): slides that survive are edited in
place, slides that carry stale numbers or a research-group-only framing are
dropped, and the missing slides are drawn with the same master, palette, and
spacing. Wording and numbers follow
`docs/research_notes/260927_labseminar_slide_outline.md` and `papers/icso/main.typ`.
"""

from __future__ import annotations

import copy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
SEMINAR = ROOT / "papers" / "seminar"
TEMPLATE = SEMINAR / "20260721_optcommrg_takamoto_v3_issl.pptx"
OUT = SEMINAR / "20260928_labseminar_takamoto.pptx"
FIG = ROOT / "papers" / "icso" / "figure"
WORK = SEMINAR / "_work" / "assets_0928"

NAVY = "082B52"
BLUE = "1B74AA"
MID_BLUE = "5BB5D8"
PALE_BLUE = "EBF6FB"
LIGHT = "F4F6F8"
BORDER = "B8D5E2"
GREEN = "239571"
PALE_GREEN = "EAF8F1"
ORANGE = "EC9427"
PALE_ORANGE = "FEF5E6"
RED = "D03B46"
PALE_RED = "FDEFF0"
PURPLE = "5D4084"
GRAY = "666F79"
DARK = "232B34"
WHITE = "FFFFFF"

FONT = "Yu Gothic"
MATH_FONT = "Cambria Math"

MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
ANS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

EMU = 914400


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------
def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def sl(prs: Presentation, number: int):
    """July deck slide by its 1-based position in the original file."""
    return prs.slides[number - 1]


def drop(shape) -> None:
    shape._element.getparent().remove(shape._element)


def by_name(slide, name: str, occurrence: int = 0):
    hits = [sh for sh in slide.shapes if sh.name == name]
    return hits[occurrence]


def by_text(slide, needle: str, occurrence: int = 0):
    hits = [
        sh
        for sh in slide.shapes
        if sh.has_text_frame and needle in sh.text_frame.text
    ]
    if not hits:
        raise KeyError(f"no shape containing {needle!r}")
    return hits[occurrence]


def set_typeface(run, name: str = FONT) -> None:
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = latin.makeelement(qn("a:ea"), {})
        latin.addnext(ea)
    ea.set("typeface", name)


def style_run(run, size=None, color=None, bold=None, font=FONT) -> None:
    set_typeface(run, font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)


def write_lines(text_frame, lines, size=None, color=None, bold=None, font=FONT,
                align=None, space_after=0.0, line_spacing=None) -> None:
    """Replace a text frame with `lines`.

    Each entry is a string or ``(text, overrides)``; overrides may set
    ``size``/``color``/``bold``/``font``/``align``/``space_before``.
    """
    text_frame.word_wrap = True
    text_frame.clear()
    base_p = text_frame.paragraphs[0]
    for index, line in enumerate(lines):
        overrides = {}
        if isinstance(line, tuple):
            line, overrides = line
        paragraph = base_p if index == 0 else text_frame.add_paragraph()
        paragraph.space_after = Pt(overrides.get("space_after", space_after))
        if overrides.get("space_before") is not None:
            paragraph.space_before = Pt(overrides["space_before"])
        if line_spacing is not None:
            paragraph.line_spacing = line_spacing
        if overrides.get("align", align) is not None:
            paragraph.alignment = overrides.get("align", align)
        run = paragraph.add_run()
        run.text = line
        style_run(
            run,
            size=overrides.get("size", size),
            color=overrides.get("color", color),
            bold=overrides.get("bold", bold),
            font=overrides.get("font", font),
        )


def retext(shape, lines, **kwargs) -> None:
    if isinstance(lines, str):
        lines = [lines]
    write_lines(shape.text_frame, lines, **kwargs)


def move(shape, x=None, y=None, w=None, h=None) -> None:
    if x is not None:
        shape.left = Inches(x)
    if y is not None:
        shape.top = Inches(y)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


def shift_below(slide, y_min: float, dy: float, y_max: float = 99.0) -> None:
    """Shift every element (including math AlternateContent) inside a band."""
    for element in slide.shapes._spTree:
        for off in element.iter(qn("a:off")):
            if off.get("y") is None:
                continue
            y = int(off.get("y")) / EMU
            if y_min <= y <= y_max:
                off.set("y", str(int(round((y + dy) * EMU))))
            break


# --------------------------------------------------------------------------
# drawing helpers, matching the July visual language
# --------------------------------------------------------------------------
def add_text(slide, text, x, y, w, h, *, size=24, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    lines = text if isinstance(text, list) else [text]
    write_lines(tf, lines, size=size, color=color, bold=bold, font=font,
                align=align, line_spacing=line_spacing)
    return shape


def add_round_rect(slide, x, y, w, h, *, fill=WHITE, line=BLUE, width=1.8,
                   shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(width)
    shape.shadow.inherit = False
    if shape.has_text_frame:
        shape.text_frame.text = ""
    return shape


def add_card(slide, title, body, x, y, w, h, *, accent=BLUE, fill=WHITE,
             title_size=27, body_size=24, centered=True, body_offset=1.05):
    add_round_rect(slide, x, y, w, h, fill=fill, line=accent)
    align = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
    if title:
        add_text(slide, title, x + 0.22, y + 0.32, w - 0.44, 0.7,
                 size=title_size, color=accent, bold=True, align=align)
    if body:
        add_text(slide, body, x + 0.30, y + body_offset, w - 0.60,
                 h - body_offset - 0.2, size=body_size, color=DARK,
                 align=align, line_spacing=1.15)


def add_metric(slide, value, label, x, y, w, h, *, accent=BLUE, fill=WHITE,
               value_size=34, label_size=19):
    add_round_rect(slide, x, y, w, h, fill=fill, line=accent)
    add_text(slide, value, x + 0.12, y + 0.18, w - 0.24, h * 0.5,
             size=value_size, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.12, y + h * 0.56, w - 0.24, h * 0.42,
             size=label_size, color=GRAY, align=PP_ALIGN.CENTER,
             line_spacing=1.05)


def style_cell(cell, text, *, size, color, bold, align, fill):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.14)
    cell.margin_top = cell.margin_bottom = Inches(0.05)
    write_lines(cell.text_frame, text.split("\n"), size=size, color=color,
                bold=bold, align=align)


def add_table(slide, rows, x, y, w, h, *, col_widths=None, size=22,
              header_size=23, header_fill=NAVY, highlight_last=False,
              first_col_align=PP_ALIGN.LEFT, body_align=PP_ALIGN.CENTER):
    graphic = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x),
                                     Inches(y), Inches(w), Inches(h))
    table = graphic.table
    tbl = graphic._element.graphic.graphicData.tbl
    tbl[0][-1].set("firstRow", "1")
    tbl[0][-1].set("bandRow", "0")
    if col_widths:
        total = sum(col_widths)
        for idx, cw in enumerate(col_widths):
            table.columns[idx].width = Inches(w * cw / total)
    last_col = len(rows[0]) - 1
    for r, row in enumerate(rows):
        table.rows[r].height = Inches(h / len(rows))
        for c, value in enumerate(row):
            if r == 0:
                style_cell(table.cell(r, c), value, size=header_size,
                           color=WHITE, bold=True,
                           align=first_col_align if c == 0 else body_align,
                           fill=header_fill)
                continue
            if highlight_last and r == len(rows) - 1:
                fill = PALE_GREEN
            elif c == last_col and last_col > 0:
                fill = PALE_BLUE
            else:
                fill = WHITE if r % 2 else LIGHT
            style_cell(table.cell(r, c), value, size=size, color=DARK,
                       bold=(highlight_last and r == len(rows) - 1),
                       align=first_col_align if c == 0 else body_align,
                       fill=fill)
    return table


def delete_table_row(table, index: int) -> None:
    tr = table._tbl.tr_lst[index]
    tr.getparent().remove(tr)


def fill_table(table, rows) -> None:
    """Rewrite cell texts, keeping the existing table formatting."""
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            paragraph = cell.text_frame.paragraphs[0]
            template = paragraph.runs[0] if paragraph.runs else None
            size = template.font.size.pt if template and template.font.size else 21
            bold = template.font.bold if template else None
            try:
                color = str(template.font.color.rgb) if template else DARK
            except (AttributeError, TypeError, ValueError):
                color = WHITE if r == 0 else DARK
            align = paragraph.alignment
            write_lines(cell.text_frame, value.split("\n"), size=size,
                        color=color, bold=bold, align=align)


def add_image(slide, path: Path, x, y, w, h):
    with Image.open(path) as image:
        ratio = image.width / image.height
    if ratio > w / h:
        dw, dh = w, w / ratio
    else:
        dh, dw = h, h * ratio
    return slide.shapes.add_picture(str(path), Inches(x + (w - dw) / 2),
                                    Inches(y + (h - dh) / 2), Inches(dw),
                                    Inches(dh))


def crop(path: Path, name: str, box) -> Path:
    WORK.mkdir(parents=True, exist_ok=True)
    out = WORK / name
    with Image.open(path) as image:
        left, top, right, bottom = box
        image.crop((round(image.width * left), round(image.height * top),
                    round(image.width * right),
                    round(image.height * bottom))).save(out)
    return out


def add_math_line(slide, segments, x, y, w, h, *, align=PP_ALIGN.CENTER):
    """One rich line; `segments` are (text, size, color, bold, baseline)."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    for text, size, color, bold, baseline in segments:
        run = paragraph.add_run()
        run.text = text
        style_run(run, size=size, color=color, bold=bold, font=MATH_FONT)
        if baseline:
            run._r.get_or_add_rPr().set("baseline", str(baseline))
    return shape


def new_slide(prs, section, title, *, title_size=32, title_lines=None):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    for placeholder in list(slide.placeholders):
        drop(placeholder)
    add_text(slide, section, 0.60, 0.24, 17.40, 0.55, size=20, color=GRAY,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, title_lines or title, 1.68, 1.42, 24.10, 1.30,
             size=title_size, color=NAVY, bold=True, valign=MSO_ANCHOR.MIDDLE,
             line_spacing=1.1)
    add_text(slide, "", 23.80, 14.02, 2.20, 0.42, size=18, color=GRAY,
             align=PP_ALIGN.RIGHT)
    return slide


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def set_section(slide, label: str) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if abs((shape.top or 0) / EMU - 0.25) < 0.10 and (shape.left or 0) / EMU < 1.0:
            retext(shape, label, size=20, color=GRAY, bold=False)
            return
    add_text(slide, label, 0.60, 0.24, 17.40, 0.55, size=20, color=GRAY,
             valign=MSO_ANCHOR.MIDDLE)


def page_number_shape(slide):
    """The page-number box: a placeholder on July 1–5, a text box afterwards."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        left = (shape.left or 0) / EMU
        top = (shape.top or 0) / EMU
        if left > 20.0 and top > 13.5 and len(shape.text_frame.text.strip()) <= 3:
            return shape
    return None


def set_title(slide, text, *, name="TextBox 2", size=32):
    """Rewrite a July title text box and give it the full title width."""
    shape = by_name(slide, name)
    retext(shape, text, size=size, color=NAVY, bold=True)
    move(shape, x=1.68, y=1.42, w=24.10, h=1.30)
    return shape


# --------------------------------------------------------------------------
# main-deck slides that are edited copies of July slides
# --------------------------------------------------------------------------
def build_title(prs):
    slide = sl(prs, 1)
    retext(by_name(slide, "タイトル 1"),
           ["衛星光通信の粗捕捉に向けた", "時変熱バイアスのフィードフォワード補正"],
           size=32, color=NAVY, bold=True)
    english = by_name(slide, "TextBox 7")
    move(english, x=10.58, y=10.95, w=14.00, h=1.10)
    retext(english,
           ["Feedforward Correction of Time-Varying Thermal Bias",
            "for Coarse Acquisition in Optical Communication Systems"],
           size=17, color=GRAY, bold=False)
    retext(by_name(slide, "字幕 2"), "高本 英熙（東京大学）", size=28,
           color=DARK, bold=True)
    retext(by_name(slide, "テキスト プレースホルダー 3"), "2026/09/28",
           size=22, color=GRAY, bold=False)
    retext(by_name(slide, "テキスト プレースホルダー 5"), "全体輪講",
           size=22, color=GRAY, bold=False)
    add_text(slide, "ICSO 2026 提出済み（10/15 口頭発表）", 16.40, 12.75, 7.80,
             0.50, size=17, color=GRAY, align=PP_ALIGN.RIGHT)
    set_notes(slide, "名乗る。今日は5月の全体輪講で話した構想の続きで、"
                     "解析と捕捉時間まで話す、とだけ言う。")
    return slide


def build_since_may(prs):
    slide = sl(prs, 5)
    set_section(slide, "前回からの更新")
    retext(by_name(slide, "タイトル 1"),
           "5月は構想と仮想バイアスだった。今日は21条件の熱構造解析と捕捉時間まで示す",
           size=31, color=NAVY, bold=True)
    since = by_name(slide, "TextBox 5")
    move(since, x=4.60, w=4.60)
    retext(since, "5/25 全体輪講", size=28, color=GRAY, bold=True,
           align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 6"), "9/28 本日", size=28, color=BLUE,
           bold=True)
    retext(
        by_name(slide, "TextBox 9"),
        [("構想と枠組みの提案", {"size": 27, "color": GRAY, "bold": True,
                                 "space_after": 14}),
         "• 熱LOSを走査中心から引く、という枠",
         "• TD/Femapを高精度参照にする提案",
         "• 軽量モデルの候補を列挙（定数・Fourier・幾何・温度）",
         "• PATの数字は仮想（mock）の熱バイアス"],
        size=25, color=DARK, bold=False, space_after=8, line_spacing=1.1)
    retext(
        by_name(slide, "TextBox 10"),
        [("解析・モデル・捕捉時間", {"size": 27, "color": BLUE, "bold": True,
                                     "space_after": 14}),
         "• 箱型衛星21ケースの温度場と熱変形を計算",
         "• 温度差と運用フラグの階層モデルを構築",
         "• 未知ケースを係数から外して予測精度を評価",
         "• 実軌道の予測誤差を足した捕捉時間を評価",
         "• 主結果は捕捉前のfeedforward。残差更新は予備"],
        size=25, color=DARK, bold=False, space_after=8, line_spacing=1.1)
    retext(by_name(slide, "TextBox 12"),
           "問題の置き方は5月と同じ。中身が構想から実解析に置き換わった",
           size=31, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    set_notes(slide, "細部を覚えている必要はない。問題の置き方は同じで、"
                     "中身が解析に置き換わった、とだけ。5月のPATの秒は仮想モデル。")
    return slide


def build_problem(prs):
    slide = sl(prs, 4)
    set_section(slide, "背景")
    retext(by_name(slide, "タイトル 1"),
           "粗捕捉では相手の光がまだ無く、走査中心のずれが探索時間を決める",
           size=32, color=NAVY, bold=True)
    order = by_name(slide, "TextBox 12")
    move(order, x=9.70, y=6.15, w=6.80, h=0.90)
    retext(order, "150 µrad〜1 mrad超の熱LOS誤差", size=25, color=RED,
           bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 7"),
           ["姿勢・軌道・熱LOSを含む", "不確定領域を探索する"],
           size=26, color=DARK, bold=False, align=PP_ALIGN.CENTER)
    heading = by_name(slide, "TextBox 15")
    move(heading, x=18.30, y=4.25, w=6.30, h=0.65)
    retext(heading, "捕捉前の予測補正", size=27, color=GREEN, bold=True,
           align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 16"),
           ["予測した熱LOSを走査中心から引き", "探索域・捕捉時間を下げる"],
           size=26, color=DARK, bold=False, align=PP_ALIGN.CENTER)
    set_notes(slide, "捕捉が終わって相手光がセンサに乗れば残差でフィードバックできる。"
                     "今日見るのはその前。熱は日照と発熱という運用条件に結びついているので、"
                     "未知の外乱として全部掃かなくてもよい可能性がある。")
    return slide


def build_satellite(prs):
    slide = sl(prs, 8)
    set_section(slide, "熱構造解析")
    set_title(slide, "評価対象は、STTとLCTを別の面に置いた箱型のLEO衛星である")
    fill_table(by_name(slide, "Table 11").table, [
        ["項目", "設定"],
        ["本体寸法", "590 × 600 × 990 mm"],
        ["パネル", "厚さ10 mmのシェル要素"],
        ["材料", "A5052"],
        ["LCT", "MZ面の中央"],
        ["STT", "PZ面の中央"],
    ])
    caption = by_text(slide, "Configuration")
    retext(caption, "解析モデルの主な設定", size=22, color=GRAY, bold=False,
           align=PP_ALIGN.CENTER)
    move(caption, x=19.04, y=3.45, w=6.96, h=0.55)
    retext(by_name(slide, "テキスト ボックス 23"), ["LOS", "視線方向 ≈ −Z"],
           size=23, color=DARK, bold=True)
    note = by_name(slide, "TextBox 12")
    move(note, x=18.40, y=12.50, w=7.60, h=1.50)
    retext(note,
           ["特定衛星の設計審査ではない。",
            "同じ構造のまま、太陽面と発熱を変える"],
           size=24, color=NAVY, bold=True, line_spacing=1.15)
    set_notes(slide, "姿勢は本体指向を想定し、通信中に太陽面が変わるケースは入れていない。"
                     "太陽面はケースごとに固定。内部発熱の代表はPROP 25 W（PY）とPCDU 10 W（MY）。"
                     "STT/LCTは通信時を想定して基本ON。")
    return slide


def build_thermal_desktop(prs):
    slide = sl(prs, 9)
    set_section(slide, "熱構造解析")
    set_title(slide, "軌道・姿勢・表面・発熱をThermal Desktopに入れ、温度場の時系列を得る")
    fill_table(by_name(slide, "Table 4").table, [
        ["入力項目", "設定"],
        ["軌道", "太陽同期 LTAN06 / 高度800 km / 冷側条件"],
        ["周期", "約6050 s（約101分）"],
        ["解析時間", "約18157 s（3周）"],
        ["サンプリング", "約60.5 s ごと、301点"],
        ["日陰", "蝕あり / 全日照"],
        ["太陽指向面", "MX / MY / PX / PY"],
        ["節点数", "125"],
    ])
    caption = by_text(slide, "TD Configuration")
    retext(caption, "熱解析の主な入力", size=22, color=GRAY, bold=False,
           align=PP_ALIGN.CENTER)
    move(caption, x=1.46, y=3.30, w=11.80, h=0.55)
    retext(by_name(slide, "TextBox 6"), "ケースで変える環境条件", size=27,
           color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 7"),
           ["• 内部機器の電源のON/OFF",
            "• パネルの表面特性：標準（α=ε=0.5）/ Black / Alodine",
            "• 軌道の種類：LTAN18 / 高度693 km（Sentinel-1衛星の模擬）"],
           size=24, color=DARK, bold=False, space_after=8, line_spacing=1.1)
    add_text(slide, "出力：各パネルの温度場 T(t)。ここは地上の解析で、オンボードでは回さない",
             1.46, 13.10, 12.00, 0.60, size=22, color=GRAY)
    set_notes(slide, "地上の解析である。食と日照で温度が軌道周期で呼吸する、"
                     "という絵は次の観察で見せる。")
    return slide


def build_femap(prs):
    slide = sl(prs, 11)
    set_section(slide, "熱構造解析")
    set_title(slide, "温度場をFemapに渡し、STTとLCTの並進と回転を得る")
    fill_table(by_name(slide, "Table 4").table, [
        ["構造解析の設定", "値"],
        ["材料", "アルミニウム A5052"],
        ["ヤング率", "70.327 GPa"],
        ["ポアソン比", "0.33"],
        ["線膨張係数", "2.376×10⁻⁵ /K"],
        ["密度", "2685 kg/m³"],
        ["基準温度", "23.9 ℃"],
        ["拘束", "STT近傍の小領域"],
    ])
    caption = by_text(slide, "Femap Configuration")
    retext(caption, "構造解析の主な設定", size=22, color=GRAY, bold=False,
           align=PP_ALIGN.CENTER)
    move(caption, x=1.31, y=3.58, w=11.90, h=0.55)
    retext(by_name(slide, "テキスト ボックス 20"), "熱解析の温度場", size=22,
           color=BLUE, bold=True)
    add_text(slide, "熱はThermal Desktop、構造はFemap。どちらも教師データを作るための地上ツール",
             1.31, 13.00, 12.50, 0.60, size=22, color=GRAY)
    set_notes(slide, "ソフトを熱と構造で分けている。軌道上の補正器そのものではない。"
                     "拘束条件の詳細は聞かれたら答える。")
    return slide


def build_los(prs):
    slide = sl(prs, 12)
    set_section(slide, "熱構造解析")
    set_title(slide, "遠方リンクで効くのは、STTに対するLCT光軸の相対回転である")
    retext(by_name(slide, "TextBox 7"), "LOS誤差の定義", size=27, color=BLUE,
           bold=True, align=PP_ALIGN.CENTER)
    move(by_name(slide, "TextBox 7"), x=17.50, y=3.55, w=7.48, h=0.62)
    bullets = by_name(slide, "TextBox 8")
    move(bullets, y=4.70, h=2.60)
    retext(
        bullets,
        ["• LCT光軸の回転 − STT姿勢基準の回転",
         "• STT自身の熱回転は姿勢の基準が回るので相殺される",
         "• 遠方通信の走査中心誤差に直接対応する",
         "• 代表点間の並進による傾きは含めない"],
        size=22, color=DARK, bold=False, space_after=8, line_spacing=1.1)
    bottom = by_name(slide, "テキスト ボックス 26")
    move(bottom, x=1.20, w=15.40)
    retext(bottom,
           "並進ではなく、遠方通信に効くSTT/LCTの相対光軸回転をLOS誤差とする",
           size=29, color=NAVY, bold=True)
    set_notes(slide, "遠方では光軸が何度回ったかが指向誤差になる。並進は診断量として残しているが"
                     "走査中心には足していない。代表節点の回転を機器姿勢の代理にしている。"
                     "取付面全体の剛体フィットはまだやっていない。")
    return slide


def build_cases(prs):
    slide = sl(prs, 13)
    set_section(slide, "熱構造解析")
    set_title(slide, "構造とLOS定義は固定し、太陽面・発熱・表面・軌道だけを変えた")
    drop(by_name(slide, "テキスト ボックス 7"))  # internal "記述先" memo
    table = by_name(slide, "Table 4").table
    delete_table_row(table, 5)  # power level
    delete_table_row(table, 5)  # orbit (merged into the environment row)
    fill_table(table, [
        ["変える条件", "ケース", "確認したいこと"],
        ["太陽指向面", "MX / MY / PX / PY", "支配軸と符号、大きさの変化"],
        ["内部発熱", "STT・LCTのみ / +PROP / +PCDU / 全発熱（半電力1ケース）",
         "ケース間の平均バイアスの変化"],
        ["表面特性", "標準（α=ε=0.5）/ Black / Alodine", "変動幅と残差の床の変化"],
        ["軌道・熱環境", "蝕あり / 全日照、LTAN06 / LTAN18",
         "熱環境・軌道条件に対する適用性"],
    ])
    move(table._graphic_frame, h=5.60)
    add_text(slide,
             ["固定：構造・寸法、STT/LCTの配置、材料、LOSの定義",
              "対象は21ケース。同じ箱の中での条件横断で、別の衛星へ係数をそのまま移す主張はしない"],
             1.20, 10.30, 24.00, 1.60, size=25, color=NAVY, bold=True,
             line_spacing=1.2)
    footnote = by_name(slide, "TextBox 6")
    move(footnote, x=1.20, y=12.40, w=24.00, h=0.55)
    retext(footnote,
           "※ PZ/MZを太陽指向面にするケースは、その面にSTT/LCTがあるので設定しない",
           size=22, color=RED, bold=False, align=PP_ALIGN.CENTER)
    set_notes(slide, "IDは読まない。予測精度は21ケース、捕捉時間は標準的なCOLD・標準表面の"
                     "14ケース、と後で分ける。")
    return slide


def build_observation(prs):
    slide = sl(prs, 15)
    set_section(slide, "解析結果")
    set_title(slide, "熱LOSは軌道周期で変わり、支配軸と平均のずれは太陽面に従う")
    left_caption = by_name(slide, "TextBox 9")
    move(left_caption, x=1.17, y=12.55, w=12.88, h=0.45)
    retext(left_caption, "Case 04（MY太陽指向）の遠方LOS", size=21, color=BLUE,
           bold=True, align=PP_ALIGN.CENTER)
    right_caption = by_name(slide, "TextBox 8")
    move(right_caption, x=14.05, y=12.55, w=12.08, h=0.45)
    retext(right_caption, "各パネルの温度", size=21, color=BLUE, bold=True,
           align=PP_ALIGN.CENTER)
    bullets = by_name(slide, "テキスト ボックス 25")
    move(bullets, x=1.20, y=13.05, w=24.20, h=1.20)
    retext(bullets,
           ["• 温度場と熱LOSは同じ軌道周期。支配軸は太陽面で決まる（MY/PY → y、MX/PX → x）",
            "• 大きさは太陽面で150〜1280 µradまで開く。表面は変動幅と残差の床、内部発熱は平均のずれに効く"],
           size=24, color=RED, bold=True, space_after=4, line_spacing=1.05)
    set_notes(slide, "だから次のモデルは、太陽面を陽に持ち、軌道の中の時間変化とケース間の平均を"
                     "分けないといけない。式はまだ出さない。MYは小さくPYは1 mrad級、とだけ。")
    return slide


def build_model_requirements(prs):
    slide = sl(prs, 17)
    set_section(slide, "軽量モデル")
    set_title(slide, "軌道上では温度差と運用フラグだけから、時変と平均を分けて予測する")
    label = by_name(slide, "TextBox 5")
    move(label, x=2.70, y=3.50, w=5.00, h=0.62)
    retext(label, "軌道上で必要なこと", size=27,
           color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 6"),
           ["• TD/Femapをオンボードで回さない",
            "• 少数の温度と運用フラグだけで計算",
            "• 係数は固定の少数",
            "• 出力は走査中心から引く角度"],
           size=25, color=DARK, bold=False, space_after=34, line_spacing=1.15)
    fill_table(by_name(slide, "Table 7").table, [
        ["候補", "長所 / 短所"],
        ["定数バイアス", "軽い / 軌道内の時間変化を表せない"],
        ["軌道位相のFourier", "周期は表せる / 軌道・条件が変わると作り直し"],
        ["取付点の温度を複数", "自由度は高い / 温度差と共線で係数が壊れる"],
        ["太陽面の温度 T", "物理的 / 時間変化を追い切れない"],
        ["太陽面の温度差 ΔT", "物理的 / ケース間の平均のずれを説明できない"],
        ["階層型 ΔT", "軌道内の時変とケース間の平均を分けて表せる"],
    ])
    retext(by_name(slide, "TextBox 8"),
           "→ 太陽面–反対面の温度差 ΔT と、ケース間の平均を分ける階層型モデルを採用",
           size=30, color=NAVY, bold=True)
    set_notes(slide, "5月に並べた候補のうち、温度差と運用フラグの組み合わせを採用した。"
                     "取付点の温度時系列を足すと、温度差とほぼ同じ動きで係数が壊れた。その失敗は裏。")
    return slide


def build_hierarchical_model(prs):
    slide = sl(prs, 19)
    level2 = sl(prs, 20)
    set_section(slide, "軽量モデル")
    set_title(slide, "軌道内の変動は太陽面ごとの感度かける温度差。平均は太陽面と発熱フラグ")

    # make room for the Level-2 expression, then copy it over from July 20
    shift_below(slide, 7.90, 0.62, 12.20)
    math_tag = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"

    def math_text(element):
        return "".join(t.text or "" for t in element.iter(math_tag))

    equation = max(
        level2.shapes._spTree.findall(MC + "AlternateContent"),
        key=lambda element: len(math_text(element)),
    )
    copied = copy.deepcopy(equation)
    for off in copied.iter(qn("a:off")):
        if off.get("y") is not None:
            off.set("y", str(int(round(7.32 * EMU))))
            break
    slide.shapes._spTree.append(copied)

    retext(by_name(slide, "TextBox 14"), "支配軸の選択", size=26, color=GREEN,
           bold=True, align=PP_ALIGN.CENTER)
    move(by_name(slide, "TextBox 14"), x=1.85, y=8.90, w=6.96, h=0.62)
    retext(by_name(slide, "TextBox 15"),
           ["太陽指向面から予測軸を選ぶ",
            "MX/PX → x軸、MY/PY → y軸",
            "非支配軸はDCだけを予測する",
            "MX/PXでは約 −600 µrad あり、",
            "ゼロにすると捕捉が落ちる"],
           size=21, color=DARK, bold=False, align=PP_ALIGN.CENTER,
           space_after=2, line_spacing=1.1)
    move(by_name(slide, "TextBox 15"), x=1.85, y=9.75, w=6.96, h=2.10)
    retext(by_name(slide, "TextBox 9"),
           ["太陽指向面ごとの熱感度",
            "同じ太陽指向面のケース間で共有",
            "+30.6 / +28.6 / −28.1 / −28.7 µrad/K"],
           size=24, color=DARK, bold=False, align=PP_ALIGN.CENTER,
           space_after=4, line_spacing=1.1)
    retext(by_name(slide, "TextBox 12"),
           ["ケース依存のDCオフセット",
            "太陽指向面と発熱のON/OFFフラグから予測",
            "下の式で分解する"],
           size=24, color=DARK, bold=False, align=PP_ALIGN.CENTER,
           space_after=4, line_spacing=1.1)
    retext(by_name(slide, "テキスト ボックス 21"),
           "時間変動は面間の温度差で、ケースごとの差はDCオフセットで表す",
           size=30, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "係数は合計16個（支配軸10＋非支配軸のDC 6）。本構造・本配置・本LOS定義のもの",
             1.40, 13.30, 24.00, 0.55, size=21, color=GRAY,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "I_prop と I_pcdu はONなら1、OFFなら0。連続のワット数や、軌道の途中で"
                     "電源が入る過渡はこの式では表していない。感度の値は他機へは再同定が必要。")
    return slide


def build_scan_definition(prs):
    slide = sl(prs, 29)
    set_section(slide, "捕捉評価の条件")
    set_title(slide, "予測を走査中心から引き、真の向きが検出円に入った時刻を捕捉とする")
    drop(by_name(slide, "Table 20"))
    drop(by_text(slide, "比較する評価条件"))
    drop(by_text(slide, "LOS予測精度ではなく"))
    predicted = sorted([sh for sh in slide.shapes
                        if sh.has_text_frame
                        and sh.text_frame.text.strip().startswith("Model-predicted")],
                       key=lambda sh: sh.top or 0)
    retext(predicted[0], "予測した熱LOS", size=21, color=WHITE, bold=True,
           align=PP_ALIGN.CENTER)
    move(predicted[0], x=1.73, y=7.95, w=4.26, h=0.55)
    retext(predicted[1], ["非熱誤差", "（軌道・姿勢など）"], size=21, color=WHITE,
           bold=True, align=PP_ALIGN.CENTER)
    move(predicted[1], x=1.73, y=9.32, w=4.26, h=0.90)
    for name, text, x, y in (
        ("TextBox 7", "熱LOSの真値", 1.71, 6.42),
        ("TextBox 13", "走査中心の残差", 8.77, 7.80),
        ("TextBox 16", "矩形スパイラル走査", 14.38, 7.80),
        ("TextBox 19", "捕捉の成否と時間", 19.92, 7.80),
    ):
        label = by_name(slide, name)
        retext(label, text, size=21, color=WHITE, bold=True,
               align=PP_ALIGN.CENTER)
        move(label, x=x, y=y, w=4.26, h=0.60)
    params = [
        ("±1600 µrad", "走査範囲", BLUE),
        ("120 µrad", "走査ステップ", BLUE),
        ("150 µrad", "検出半径（ビーコン級）", GREEN),
        ("0.1 s / 点", "滞在時間", GREEN),
        ("27 × 27 = 729点", "最大走査点数", GREEN),
    ]
    for index, (value, label, accent) in enumerate(params):
        add_metric(slide, value, label, 1.35 + index * 4.85, 10.45, 4.35, 1.95,
                   accent=accent, value_size=26, label_size=18)
    add_text(slide, "格子対角の半分 ≈ 85 µrad < 検出半径 150 µrad なので、被覆の穴はない",
             1.35, 12.62, 23.50, 0.55, size=25, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "点間移動・整定時間・確率的な検出は無視。秒はこの走査条件での比較用",
             1.35, 13.22, 23.50, 0.55, size=22, color=RED,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "粗捕捉は通信ビームではなく、拡がったビーコンで掃く。検出半径150 µradは"
                     "0.3 mrad級ビーコンの半角のオーダー。5月の2秒前後と、途中で一度出した"
                     "100秒超は、どちらも今日の数字ではない。前者は仮想モデル、後者は"
                     "通信ビーム級の密な掃引だった。")
    return slide


def build_nonthermal(prs):
    slide = sl(prs, 31)
    set_section(slide, "捕捉評価の条件")
    set_title(slide, "熱のほかに、軌道予測・アライメント・姿勢・ドリフトを足す")
    fill_table(by_name(slide, "Table 6").table, [
        ["成分", "作り方", "時間特性", "大きさ"],
        ["軌道予測", "Sentinel-1の最新TLEを伝播し、高精度暦POEORBを真値として横断面へ射影",
         "軌道周期", "平均 280–700 µrad"],
        ["アライメント", "ケースごとに一定の2軸残差", "DC", "50 µrad/軸, 1σ"],
        ["姿勢", "サンプルごとの正規乱数", "簡易な広帯域", "50 µrad/軸, 1σ"],
        ["ドリフト", "正弦波", "周期 900 s", "振幅 30 µrad"],
    ])
    add_text(slide, "ケースごとに乱数は1本。統計の幅は見ていない",
             1.20, 11.20, 24.20, 0.55, size=22, color=GRAY)
    set_notes(slide, "熱を引いたあとに残る床が、だいたいこの非熱の大きさになる。"
                     "射影の座標の話は、聞かれたら裏の図で示す。")
    return slide


def build_thermal_only(prs):
    slide = sl(prs, 35)
    set_section(slide, "捕捉評価の結果")
    set_title(slide, "熱だけなら平均捕捉時間は14.6秒から0.10秒になり、熱真値の補正と同じになる",
              size=31)
    for name in ("TextBox 8", "Rectangle 9", "Rectangle 10", "TextBox 11"):
        drop(by_name(slide, name))
    for name, dy in (("TextBox 12", -1.62), ("Rectangle 13", -1.62),
                     ("Rectangle 14", -1.62), ("TextBox 15", -1.62),
                     ("TextBox 16", -1.62), ("Rectangle 17", -1.62),
                     ("Rectangle 18", -1.62), ("TextBox 19", -1.62)):
        shape = by_name(slide, name)
        move(shape, y=(shape.top / EMU) + dy)
    retext(by_name(slide, "TextBox 4"), "補正なし", size=23, color=DARK,
           bold=False)
    retext(by_name(slide, "TextBox 12"), "階層モデル", size=23, color=DARK,
           bold=False)
    retext(by_name(slide, "TextBox 16"), "熱の真値を引いた上界", size=23,
           color=DARK, bold=False)
    retext(by_name(slide, "TextBox 7"), "14.6 s", size=22, color=RED,
           bold=True)
    retext(by_name(slide, "TextBox 15"), "0.10 s", size=22, color=GREEN,
           bold=True)
    retext(by_name(slide, "TextBox 19"), "0.10 s", size=22, color=BLUE,
           bold=True)
    move(by_name(slide, "Rectangle 14"), w=0.10)
    move(by_name(slide, "Rectangle 18"), w=0.10)
    table = by_name(slide, "Table 20").table
    delete_table_row(table, 2)
    fill_table(table, [
        ["補正方法", "成功率", "平均捕捉時間"],
        ["補正なし", "100%", "14.6 s"],
        ["階層モデル", "100%", "0.10 s"],
        ["熱の真値", "100%", "0.10 s"],
    ])
    retext(by_name(slide, "TextBox 22"), "9.3 µrad", size=34, color=GREEN,
           bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 23"), "補正後の平均熱残差", size=20,
           color=GRAY, bold=False, align=PP_ALIGN.CENTER)
    residual = by_name(slide, "TextBox 24")
    move(residual, x=2.20, y=11.40, w=22.20, h=1.30)
    retext(residual,
           ["残差が検出半径150 µradより十分小さいので、滞在1回分で捕捉する",
            "COLD・標準表面の14ケース平均。主結果は次の非熱込み"],
           size=25, color=GRAY, bold=False, align=PP_ALIGN.CENTER,
           line_spacing=1.15)
    set_notes(slide, "これは熱モデル単体の能力。真値と同じ秒なのは、残差が検出円の中に"
                     "入ったからで、波形が完全に一致したからではない。")
    return slide


def build_with_nonthermal(prs):
    slide = sl(prs, 36)
    set_section(slide, "捕捉評価の結果")
    set_title(slide, "非熱を足すと19.2秒から5.45秒。残る探索は熱ではなく非熱の床である",
              size=31)
    retext(by_name(slide, "TextBox 4"), "補正なし", size=23, color=DARK,
           bold=False)
    retext(by_name(slide, "TextBox 8"), "階層モデル", size=23, color=DARK,
           bold=False)
    retext(by_name(slide, "TextBox 7"), "19.2 s", size=22, color=RED,
           bold=True)
    retext(by_name(slide, "TextBox 11"), "5.45 s", size=22, color=GREEN,
           bold=True)
    move(by_name(slide, "Rectangle 10"), w=9.34 * 5.45 / 19.2)
    retext(by_name(slide, "TextBox 13"), "−72%", size=34, color=GREEN,
           bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 14"), "平均捕捉時間", size=20, color=GRAY,
           bold=False, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 16"), "98 → 100%", size=32, color=BLUE,
           bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 17"), "成功率", size=20, color=GRAY,
           bold=False, align=PP_ALIGN.CENTER)
    fill_table(by_name(slide, "Table 18").table, [
        ["指標", "補正なし", "階層モデル（フィードフォワード）"],
        ["成功率", "98%", "100%"],
        ["平均捕捉時間", "19.2 s", "5.45 s"],
        ["補正後の平均初期誤差", "—", "448 µrad（ほぼ非熱）"],
        ["PY太陽面（熱 約1.2 mrad）", "37.3–39.5 s / 88–97%", "1.3–1.8 s / 100%"],
    ])
    retext(by_name(slide, "TextBox 19"),
           "MY（熱150–260 µrad）の改善は小さい。大きい時変の熱を落とし、探索を非熱の床まで下げる",
           size=26, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    move(by_name(slide, "TextBox 19"), x=1.50, w=23.70)
    set_notes(slide, "手法の仕事は全誤差を消すことではない。平均5.45秒は熱が大きい面が"
                     "引き上げた結果でもある。平均だけを引いた場合の新しい走査での秒は、"
                     "まだ出していない。")
    return slide


# --------------------------------------------------------------------------
# main-deck slides drawn from scratch
# --------------------------------------------------------------------------
def build_error_budget(prs):
    slide = new_slide(prs, "背景",
                      "熱は軌道予測誤差と同じかそれ以上になり得る。捕捉前は光では分離できない",
                      title_size=31)
    add_table(slide, [
        ["誤差源", "大きさの目安", "時間変化"],
        ["熱LOS（本研究の解析値）", "150 µrad – 1.3 mrad", "軌道周期（約100分）"],
        ["軌道予測（TLE伝播 vs 高精度暦）", "数百 µrad", "軌道周期（約100分）"],
        ["姿勢・較正済みアライメント", "数十 µrad 級", "DC〜広帯域"],
        ["機械振動", "数〜数十 µrad", "高周波（捕捉後の精追尾側）"],
    ], 1.60, 3.50, 23.40, 5.60, col_widths=[1.5, 1.0, 1.1], size=24,
        header_size=24)
    add_text(slide, "※ 熱のみ本研究の解析値。ほかは文献と一般的な値のオーダー",
             1.60, 9.25, 23.40, 0.50, size=20, color=GRAY)
    add_round_rect(slide, 1.60, 10.30, 23.40, 2.10, fill=PALE_BLUE,
                   line=MID_BLUE)
    add_text(slide,
             ["熱と軌道予測は、どちらも約100分の軌道周期をもつ",
              "捕捉後の残差を周波数で分けても、熱だけを取り出すことはできない"],
             2.00, 10.55, 22.60, 1.60, size=27, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.2)
    add_text(slide, "→ 観測から熱を分離するのではなく、温度と運用状態という別の入力で、捕捉の前に熱の分だけ引く",
             1.60, 12.75, 23.40, 0.60, size=24, color=GREEN, bold=True,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "5月に『軌道予測誤差も同じ帯域では』と指摘された点はその通りだった。"
                     "だから本研究は観測から熱を分離する話ではない。TLEの不連続や射影の手順は裏。")
    return slide


def build_scope(prs):
    slide = new_slide(prs, "背景",
                      "予測できる熱成分だけを走査中心から引く。全指向誤差の除去ではない",
                      title_size=32)
    add_round_rect(slide, 1.60, 3.40, 23.40, 1.80, fill=PALE_GREEN, line=GREEN)
    add_text(slide, "温度と運用状態から熱LOSを予測し、粗捕捉の走査中心を捕捉前に補正する",
             1.90, 3.75, 22.80, 1.10, size=31, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
    add_table(slide, [
        ["立場", "熱の扱いと、評価している出力"],
        ["光通信のPAT", "熱は走査範囲や構造設計で吸収することが多い。捕捉時間は評価されている"],
        ["地球観測・JANUS などの温度補正",
         "温度差とLOSの一次関係は示されている。粗捕捉時間は評価されていない"],
        ["本研究",
         "バス上で離れたSTTとLCTの相対LOSを予測し、粗捕捉時間まで評価する"],
    ], 1.60, 5.90, 23.40, 4.80, col_widths=[1.0, 2.5], size=24,
        header_size=24, highlight_last=True, body_align=PP_ALIGN.LEFT)
    add_text(slide,
             ["一次式そのものが新しいのではない。対象が光学ヘッドの中ではなく衛星バス上の2機器であること、",
              "出力が画像の指向ではなく捕捉時間であることが違い"],
             1.60, 11.25, 23.40, 1.30, size=25, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.2)
    add_text(slide, "全部の誤差を消す研究ではない。あとで非熱が残っても、想定どおりである",
             1.60, 12.75, 23.40, 0.60, size=24, color=RED, bold=True,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "関連研究の詳細表は裏B8。ここでは、全部の誤差を消す研究ではない、"
                     "と先に言っておく。")
    return slide


def build_identification(prs):
    slide = new_slide(prs, "軽量モデル",
                      "テストするケースは、感度にも平均の係数にも使わない",
                      title_size=32)
    steps = [
        (1.73, 6.77, GREEN, "① ケースごとに当てはめる",
         "各ケースの最初の1周を使い、\nそのケースだけの感度 a と\n平均 b を推定する"),
        (9.25, 7.28, BLUE, "② 太陽面ごとに共有する",
         "同じ太陽指向面の a の中央値を\n共有感度にする\n\n平均 b は、太陽面と発熱フラグから\n予測する係数を、ケースを横断して決める"),
        (17.26, 7.67, ORANGE, "③ 未知ケースで試す",
         "nested leave-one-case-out\n\n評価する1ケースを、共有感度と\n両軸の平均モデルの両方から外す\n\nそのケースの後続2軌道で誤差を測る"),
    ]
    for x, w, accent, head, body in steps:
        add_card(slide, head, body, x, 3.70, w, 6.40, accent=accent,
                 fill=WHITE, title_size=27, body_size=24, body_offset=1.35)
    for x in (8.72, 16.74):
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),
                                       Inches(6.55), Inches(0.42), Inches(0.72))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = rgb(BLUE)
        arrow.line.fill.background()
        arrow.shadow.inherit = False
    add_round_rect(slide, 2.23, 10.80, 22.20, 1.35, fill=PALE_RED, line=RED)
    add_text(slide, "当たっているように見えるのは、そのケースの答えを係数に入れたから、ではない",
             2.60, 11.10, 21.50, 0.70, size=28, color=RED, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "同一の解析モデルの中での条件横断であり、別構造への検証ではない",
             2.23, 12.35, 22.20, 0.50, size=21, color=GRAY,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "名前は長いので、口頭では『未知の条件を1本外して当てる』でよい。"
                     "同一の解析モデルの中での話である、とここで一度言う。")
    return slide


def build_accuracy(prs):
    slide = new_slide(prs, "軽量モデル",
                      "未知ケースでも支配軸の誤差は中央値4.9 µrad。生の615 µradから1–2桁落ちる",
                      title_size=31)
    add_round_rect(slide, 1.10, 3.30, 12.20, 6.90, fill=None, line=BORDER,
                   width=1.0, shape_type=MSO_SHAPE.RECTANGLE)
    add_round_rect(slide, 13.55, 3.30, 12.20, 6.90, fill=None, line=BORDER,
                   width=1.0, shape_type=MSO_SHAPE.RECTANGLE)
    add_image(slide, FIG / "p3_a_emp_by_sunface.png", 1.25, 3.45, 11.90, 6.60)
    add_image(slide, FIG / "p3_b_emp_vs_b_pred.png", 13.70, 3.45, 11.90, 6.60)
    add_text(slide, "太陽指向面ごとの熱感度（ケース別と共有値）", 1.10, 10.30,
             12.20, 0.50, size=21, color=BLUE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "ケース平均DC：経験値と予測値の対比", 13.55, 10.30, 12.20,
             0.50, size=21, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_metric(slide, "615 µrad", "生LOS 支配軸RMSの中央値", 1.60, 11.00, 7.20,
               1.95, accent=BLUE, value_size=32)
    add_metric(slide, "4.9 µrad", "未知ケースのtest RMSE 中央値（平均5.5）",
               9.70, 11.00, 7.20, 1.95, accent=GREEN, value_size=32,
               label_size=18)
    add_metric(slide, "3.8 µrad", "ケース平均DCのLOO RMSE", 17.80, 11.00, 7.20,
               1.95, accent=ORANGE, value_size=32)
    add_text(slide, "共有感度：MX +30.6 / MY +28.6 / PX −28.1 / PY −28.7 µrad/K",
             1.60, 13.20, 23.40, 0.50, size=22, color=GRAY,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "標準のCOLDではおおむね3–7 µrad。悪いのはBlack（約13）、HOT（DCが数µrad残る）、"
                     "PROP半電力（約16。ON/OFFフラグでは25 Wと12.5 Wを区別できない）。詳細は限界の枚と裏。")
    return slide


def build_case08(prs):
    plot = crop(FIG / "p2_bcase_true_vs_pred_case08.png",
                "case08_no_log_title.png", (0.0, 0.055, 1.0, 0.99))
    slide = new_slide(prs, "軽量モデル",
                      "熱が最大級のCase 08でも、1250 µradが3.9 µradまで落ちる",
                      title_size=32)
    add_round_rect(slide, 1.10, 3.30, 15.40, 9.60, fill=None, line=BORDER,
                   width=1.0, shape_type=MSO_SHAPE.RECTANGLE)
    add_image(slide, plot, 1.30, 3.45, 15.00, 9.30)
    add_text(slide, "Case 08（PY太陽指向・全発熱）", 1.10, 12.95, 15.40, 0.50,
             size=22, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_metric(slide, "1250 µrad", "生の支配軸RMS", 17.20, 3.80, 8.20, 2.30,
               accent=RED, fill=PALE_RED, value_size=36)
    add_metric(slide, "3.9 µrad", "未知ケースでのtest RMSE", 17.20, 6.40, 8.20,
               2.30, accent=GREEN, fill=PALE_GREEN, value_size=36)
    add_card(slide, "この残差なら1点目で当たる",
             "補正後の残差は検出半径 150 µrad より\n十分に小さい",
             17.20, 9.00, 8.20, 2.60, accent=BLUE, fill=WHITE, title_size=25,
             body_size=23, body_offset=1.20)
    add_text(slide, "軌道の中の山を、温度差が追っている", 17.20, 11.85, 8.20,
             0.60, size=24, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    set_notes(slide, "半電力の例はこの枚に混ぜない（裏B6）。")
    return slide


def build_residual_update(prs):
    slide = new_slide(prs, "捕捉評価の結果",
                      "初回を救うのは温度の予測。次の周の周期的な床は、捕捉後の残差で下げられる",
                      title_size=31)
    add_round_rect(slide, 1.10, 3.30, 15.00, 7.60, fill=None, line=BORDER,
                   width=1.0, shape_type=MSO_SHAPE.RECTANGLE)
    add_image(slide, FIG / "fig_residual_update_comparison.png", 1.25, 3.45,
              14.70, 7.30)
    add_table(slide, [
        ["Case 16（PY）", "初周", "周1以降"],
        ["全誤差をまとめてFourier", "39.5 s", "0.38 s"],
        ["階層モデル＋残差Fourier", "1.47 s", "0.40 s"],
    ], 16.70, 3.60, 8.90, 3.30, col_widths=[1.9, 0.8, 0.9], size=21,
        header_size=21)
    add_card(slide, "定数の残差更新では足りない",
             "Case 13 の周1以降は 1.66 → 1.35 s までで、\nFourier の 0.38 s には届かない",
             16.70, 7.30, 8.90, 2.70, accent=ORANGE, fill=PALE_ORANGE,
             title_size=24, body_size=22, body_offset=1.15)
    add_card(slide, "この節の身分",
             "熱モデルの係数は更新していない。\n捕捉が成功したあとの全指向残差に対する、\n予備的な数値実験（2ケース・密サンプル）",
             16.70, 10.30, 8.90, 3.00, accent=GRAY, fill=LIGHT, title_size=24,
             body_size=21, body_offset=1.15)
    add_text(slide,
             ["主結果は温度と運用フラグのfeedforwardのまま。",
              "繰り返し通信する相手には、床をさらに下げられるという予備的な見通し"],
             1.10, 11.20, 15.00, 1.40, size=24, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.25)
    set_notes(slide, "5月に構想したAdaptiveの着地はここまで。光通信の初回捕捉は、前の周の同じ位相の"
                     "残差が無いことが多い。60秒刻みの全点を使っており、実センサの疎な成功点だけでは"
                     "まだ見ていない。")
    return slide


def build_limitations(prs):
    slide = new_slide(prs, "まとめ",
                      "同じ解析モデルの中の数値評価であり、いくつかの床が残っている",
                      title_size=32)
    items = [
        ("真値と係数が同じ解析から出ている",
         "Thermal Desktop / Femap は共通。外しているのは\nケースであって、構造やLOS定義ではない",
         BLUE, PALE_BLUE),
        ("LOSは代表節点の回転",
         "取付面の剛体フィットは未実施。地上試験も無い",
         ORANGE, PALE_ORANGE),
        ("条件が外れると残差が上がる",
         "Black 約13 µrad、HOT はDCが数µrad、半電力 約16 µrad。\n被覆・軌道・連続電力は平均モデルに未反映",
         RED, PALE_RED),
        ("捕捉評価の簡略化",
         "非熱は1本の乱数。走査の移動時間は無視。\n残差Fourier は2ケースの密サンプル",
         GREEN, PALE_GREEN),
    ]
    for index, (head, body, accent, fill) in enumerate(items):
        x = 1.40 + (index % 2) * 12.40
        y = 3.50 + (index // 2) * 4.20
        add_card(slide, head, body, x, y, 11.70, 3.70, accent=accent,
                 fill=fill, title_size=26, body_size=22, centered=False,
                 body_offset=1.25)
    add_round_rect(slide, 2.60, 12.10, 21.50, 1.30, fill=WHITE, line=NAVY)
    add_text(slide, "感度 約30 µrad/K は本構造のもの。温度差を0.1 K読み違えると約3 µrad",
             2.90, 12.40, 20.90, 0.70, size=26, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "他の衛星へは式の形を試し、係数は取り直す。Constant-bias only が未計算である"
                     "ことは、聞かれたら裏B7で答える。")
    return slide


def build_summary(prs):
    slide = new_slide(prs, "まとめ",
                      "予測できる熱を捕捉前に除くと、残る探索は非熱誤差で決まる",
                      title_size=32)
    add_metric(slide, "4.9 µrad", "未知ケースの支配軸RMSE 中央値\n（生の中央値 615 µrad）",
               1.60, 3.70, 7.40, 3.10, accent=BLUE, fill=PALE_BLUE,
               value_size=40, label_size=21)
    add_metric(slide, "14.6 → 0.10 s", "熱のみの平均捕捉時間", 9.60, 3.70, 7.40,
               3.10, accent=GREEN, fill=PALE_GREEN, value_size=36,
               label_size=21)
    add_metric(slide, "19.2 → 5.45 s", "非熱込みの平均捕捉時間\n成功率 98% → 100%",
               17.60, 3.70, 7.40, 3.10, accent=ORANGE, fill=PALE_ORANGE,
               value_size=36, label_size=21)
    add_round_rect(slide, 1.60, 7.60, 23.40, 3.20, fill=WHITE, line=NAVY)
    add_text(slide,
             ["評価した箱型衛星では、光フィードバックの前のリンク確立を短くし得る",
              "新しいのは温度と角度の一次式そのものではなく、バス上のSTT–LCTに適用し、",
              "係数を条件横断で共有し、捕捉時間まで繋いだこと"],
             2.00, 8.00, 22.60, 2.50, size=27, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.25)
    add_text(slide, "ICSO 2026 で 10月15日に口頭発表", 1.60, 11.35, 23.40, 0.60,
             size=24, color=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "ご清聴ありがとうございました", 1.60, 12.30, 23.40, 0.80,
             size=30, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    set_notes(slide, "質疑へ。")
    return slide


# --------------------------------------------------------------------------
# backup slides
# --------------------------------------------------------------------------
def build_backup_divider(prs):
    slide = sl(prs, 42)
    retext(by_text(slide, "Appendix"), "補足（質疑用）", size=44, color=WHITE,
           bold=True)
    retext(by_name(slide, "TextBox 1"), "B", size=32, color=MID_BLUE, bold=True)
    return slide


def build_b1(prs):
    slide = sl(prs, 45)
    set_section(slide, "補足")
    retext(by_name(slide, "TextBox 2"), "B1｜21ケースの内訳と評価ID",
           size=32, color=NAVY, bold=True)
    fill_table(by_name(slide, "Table 4").table, [
        ["ケース", "太陽指向面", "発熱・条件"],
        ["04, 13–15", "MY", "全発熱 / +PROP / +PCDU / STT・LCTのみ"],
        ["08, 16, 18, 19", "PY", "全発熱 / STT・LCTのみ / +PROP / +PCDU"],
        ["05, 06, 20, 21", "PX", "全発熱 / STT・LCTのみ / +PROP / +PCDU"],
        ["09, 17, 23, 24", "MX", "全発熱 / STT・LCTのみ / +PROP / +PCDU"],
        ["10–12", "MY", "全日照 / 太陽面Black / 全面Alodine"],
        ["22", "MY", "PROP 12.5 W（半電力）"],
        ["25", "MY", "LTAN18 / 高度693 km"],
    ])
    add_text(slide,
             ["評価IDは 04–06 と 08–25。01–03 と 07 はMZ太陽や設定確認のケースで除外",
              "PZ/MZを太陽面にするケースは、その面にSTT/LCTがあるので設定しない",
              "捕捉時間の見出しは、COLD・標準表面の14ケース（04–06, 08–09, 13–21）"],
             1.40, 11.40, 23.80, 2.00, size=23, color=GRAY, line_spacing=1.25)
    return slide


def build_b2(prs):
    slide = sl(prs, 30)
    set_section(slide, "補足")
    retext(by_name(slide, "TextBox 2"), "B2｜矩形スパイラル走査の幾何",
           size=32, color=NAVY, bold=True)
    fill_table(by_name(slide, "Table 27").table, [
        ["パラメータ", "設定値"],
        ["走査範囲", "±1600 µrad"],
        ["点間隔", "120 µrad"],
        ["検出半径", "150 µrad"],
        ["滞在時間", "0.1 s / 点"],
        ["総走査点数", "27 × 27 = 729点"],
        ["最大走査時間", "72.9 s"],
    ])
    drop(by_text(slide, "後で再検討"))
    add_text(slide, "格子対角の半分 120/√2 ≈ 85 µrad < 検出半径 150 µrad なので掃き残しは無い",
             13.40, 9.90, 11.90, 1.00, size=23, color=GREEN, bold=True,
             line_spacing=1.2)
    return slide


def build_b3(prs):
    slide = sl(prs, 33)
    set_section(slide, "補足")
    retext(by_name(slide, "TextBox 2"),
           "B3｜軌道予測誤差の作り方と、リンク横断面への射影",
           size=32, color=NAVY, bold=True)
    retext(by_name(slide, "TextBox 7"),
           ["• 真値は Sentinel-1 の高精度暦 POEORB",
            "• 予測は直近のTLEを SGP4 で前方伝播",
            "• 差をリンク横断面へ射影し、x–y の2成分にする",
            "• 熱LOSと同じ2成分で足し合わせる",
            "• 時系列の不連続は、TLEが更新された時刻"],
           size=24, color=DARK, bold=False, space_after=10, line_spacing=1.1)
    fill_table(by_name(slide, "Table 8").table, [
        ["太陽指向面", "通信相手とリンク", "扱い"],
        ["MY", "進行方向の衛星（高度800 km）", "評価に使う"],
        ["PY", "逆進行方向の衛星（高度800 km）", "評価に使う"],
        ["PX", "直下の地上局（約695 km）", "評価に使う"],
        ["MX", "天頂側に置いた仮想の相手（800 km）", "幾何のみ・非現実的"],
    ])
    return slide


def build_b4(prs):
    slide = sl(prs, 34)
    set_section(slide, "補足")
    retext(by_name(slide, "TextBox 2"),
           "B4｜熱と軌道予測は同じ軌道周期。周波数では分けない",
           size=32, color=NAVY, bold=True)
    fill_table(by_name(slide, "Table 6").table, [
        ["リンク", "位置誤差の平均"],
        ["MY：進行方向の衛星", "約285 µrad"],
        ["PY：逆方向の衛星", "約285 µrad"],
        ["PX：直下の地上局", "約694 µrad"],
        ["MX：天頂側の仮想", "約612 µrad"],
    ])
    retext(by_name(slide, "TextBox 8"),
           ["熱：約101分", "軌道予測：約101分 と約50分"],
           size=24, color=ORANGE, bold=True, align=PP_ALIGN.CENTER,
           line_spacing=1.2)
    retext(by_name(slide, "TextBox 9"),
           ["熱ひずみと軌道予測誤差は、大きさも周波数帯も同程度",
            "『低周波＝熱』という分離は主張しない。温度という別入力で捕捉前に引く"],
           size=25, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
           line_spacing=1.2)
    move(by_name(slide, "TextBox 9"), x=1.20, y=12.30, w=24.20, h=1.30)
    return slide


def build_b5(prs):
    slide = sl(prs, 47)
    set_section(slide, "補足")
    retext(by_name(slide, "TextBox 2"), "B5｜LOS定義の比較と、並進を入れない理由",
           size=32, color=NAVY, bold=True)
    add_text(slide,
             ["代表節点の回転を機器姿勢の代理にしている。取付面全体の剛体フィットは今後",
              "代表点間の並進による傾きは、衛星間距離で割ると小さいので捕捉評価には入れていない"],
             1.20, 12.40, 24.00, 1.30, size=23, color=GRAY, line_spacing=1.25)
    return slide


def build_b6(prs):
    slide = new_slide(prs, "補足", "B6（1/2）｜モデルが難しいケース", title_size=32)
    add_table(slide, [
        ["ケース", "残差", "何が表せていないか"],
        ["太陽面Black被覆", "約13 µrad",
         "被覆を変えると感度と変動幅が動く。係数は標準表面で決めている"],
        ["PROP半電力（12.5 W）", "約16 µrad",
         "ON/OFFの二値フラグでは25 Wと12.5 Wを区別できない"],
        ["全日照（HOT）軌道", "DCが数µrad",
         "蝕がなく熱がほぼ一定。平均の予測が軌道条件の外へ出ている"],
    ], 1.60, 3.60, 23.40, 5.20, col_widths=[1.0, 0.6, 2.4], size=24,
        header_size=24)
    add_round_rect(slide, 1.60, 9.60, 11.40, 3.20, fill=PALE_BLUE,
                   line=MID_BLUE)
    add_text(slide, "共通する構図", 2.00, 9.85, 10.60, 0.55, size=26,
             color=NAVY, bold=True)
    add_text(slide,
             ["時変の項ではなく、ケース平均の側が外れている。",
              "被覆・軌道・連続発熱量はいまの入力に入っていない"],
             2.00, 10.55, 10.60, 2.00, size=24, color=DARK, line_spacing=1.25)
    add_round_rect(slide, 13.60, 9.60, 11.40, 3.20, fill=PALE_GREEN,
                   line=GREEN)
    add_text(slide, "主結果の範囲", 14.00, 9.85, 10.60, 0.55, size=26,
             color=GREEN, bold=True)
    add_text(slide,
             ["標準COLD・標準表面の条件で、固定16係数のまま",
              "1〜2桁の低減。被覆・軌道・連続量の明示は今後"],
             14.00, 10.55, 10.60, 2.00, size=24, color=DARK, line_spacing=1.25)
    set_notes(slide, "被覆や半電力を聞かれたときの枚。発熱の係数を太陽面に依らず共有して"
                     "いる点は、まだ検証していないと正直に言う。")
    return slide


def build_b6b(prs):
    slide = sl(prs, 43)
    set_section(slide, "補足")
    retext(by_name(slide, "TextBox 2"),
           "B6（2/2）｜取付点の温度を足すと係数が壊れる",
           size=32, color=NAVY, bold=True)
    retext(by_name(slide, "TextBox 5"), "試したこと", size=27, color=NAVY,
           bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 7"),
           ["• 温度差と取付点温度の相関が 0.995",
            "• 取付点まわりの温度のばらつきは 0.1 K 程度",
            "• 係数が巨大化し、ケース間で不安定になる"],
           size=25, color=DARK, bold=False, space_after=12, line_spacing=1.1)
    retext(by_name(slide, "TextBox 10"), "解釈", size=27, color=NAVY,
           bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 11"),
           ["効いていたのは軌道内の変動ではなく", "ケース平均のDCだった"],
           size=26, color=DARK, bold=False, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 12"), "→ b_case 側へ分離した", size=27,
           color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    retext(by_name(slide, "TextBox 13"),
           "RMSEだけでなく、係数の安定性と運用上の解釈でモデルを選んだ",
           size=28, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    return slide


def build_b7(prs):
    slide = new_slide(prs, "補足",
                      "B7｜平均だけを引いた場合（定数バイアスのみ）は未計算",
                      title_size=32)
    add_round_rect(slide, 3.20, 4.20, 20.20, 4.60, fill=LIGHT, line=BORDER)
    add_text(slide, "未計算", 3.20, 4.90, 20.20, 1.20, size=48, color=GRAY,
             bold=True, align=PP_ALIGN.CENTER)
    add_text(slide,
             ["ΔT(t) を 0 とし、予測したケース平均DCだけを引いた場合の捕捉時間は、",
              "今回の走査条件ではまだ計算していない"],
             3.20, 6.40, 20.20, 1.80, size=27, color=DARK,
             align=PP_ALIGN.CENTER, line_spacing=1.25)
    add_card(slide, "この列が示すはずのこと",
             "改善のどれだけが平均の除去で、\nどれだけが軌道内の時変項によるのかを分ける",
             3.20, 9.30, 9.70, 2.80, accent=BLUE, fill=WHITE, title_size=25,
             body_size=22, body_offset=1.20)
    add_card(slide, "7月に出した秒は流用しない",
             "あのときの走査は点間隔と検出半径が\n通信ビーム級に細かく、今日の条件とは別",
             13.70, 9.30, 9.70, 2.80, accent=RED, fill=PALE_RED,
             title_size=25, body_size=22, body_offset=1.20)
    add_text(slide, "本編の結論には、この未計算の列の数字を一切使っていない",
             3.20, 12.50, 20.20, 0.60, size=24, color=RED, bold=True,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "聞かれたら、平均だけを引いた列は未計算だと答える。7月の静的バイアス"
                     "補正の秒は、点間隔40 µrad・検出半径25 µradの旧走査のもので、"
                     "今日の条件では比較にならない。")
    return slide


def build_b8(prs):
    slide = sl(prs, 44)
    set_section(slide, "補足")
    retext(by_name(slide, "TextBox 2"), "B8｜先行研究の位置づけ（2軸で見る）",
           size=32, color=NAVY, bold=True)
    fill_table(by_name(slide, "Table 4").table, [
        ["研究", "熱状態からLOSを予測するか / 粗捕捉時間を評価するか"],
        ["Riesing 2023（TBIRD）", "熱は明示せず姿勢系で吸収 / 軌道上で捕捉を評価"],
        ["Shi 2023", "構造設計で熱変形を小さくする / 捕捉時間を評価"],
        ["Hu 2022（GEO）", "周期モデルで熱LOSを補正 / 粗捕捉は評価しない"],
        ["Li 2025（LEO）", "ニューラルネットで熱LOSを補正 / 粗捕捉は評価しない"],
        ["JANUS（Turella）", "壁面温度差に比例する一次式 / 粗捕捉は評価しない"],
        ["本研究", "太陽面温度差と運用フラグで予測する / 捕捉時間まで評価する"],
    ])
    add_text(slide,
             "新しいのは一次関係そのものではなく、バス上のSTT–LCTへの適用、条件横断の係数共有、捕捉時間への接続",
             1.40, 11.70, 23.80, 0.70, size=24, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "何が新しいのかを聞かれたときの枚。予測するか、捕捉時間まで見るか、"
                     "の2軸で置く。")
    return slide


def build_b9(prs):
    slide = new_slide(prs, "補足", "B9｜解析モデルの妥当性をどう検証するか",
                      title_size=32)
    add_card(slide, "今日の主張の範囲",
             "同一の解析チェーン（Thermal Desktop / Femap）の中での数値評価である。\n"
             "外しているのはケースであって、構造やLOS定義ではない",
             1.60, 3.60, 11.60, 3.90, accent=BLUE, fill=PALE_BLUE,
             title_size=27, body_size=23, centered=False, body_offset=1.25)
    add_card(slide, "入っていないもの",
             "打上げ後のアライメント変化、実機の剛性差、組立公差、\n軌道上での経年変化",
             14.00, 3.60, 11.40, 3.90, accent=ORANGE, fill=PALE_ORANGE,
             title_size=27, body_size=23, centered=False, body_offset=1.25)
    add_card(slide, "今後の検証",
             "地上の熱真空試験で温度と変位を同時に測り、感度 a を実測で確かめる。\n"
             "係数の取り直しが前提であり、式の形が移せるかを見る",
             1.60, 8.00, 23.80, 3.40, accent=GREEN, fill=PALE_GREEN,
             title_size=27, body_size=24, centered=False, body_offset=1.25)
    add_text(slide, "主張を『解析を真値として完璧に当てる』に置いていない",
             1.60, 11.90, 23.80, 0.70, size=26, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "5月に出たモデルの妥当性の問いへの答え。同じ解析の中の評価であることを"
                     "認めた上で、地上試験で感度を確かめる筋を示す。")
    return slide


def build_b10(prs):
    slide = new_slide(prs, "補足", "B10｜姿勢の前提と、変わる頻度",
                      title_size=32)
    add_card(slide, "前提",
             "本体指向を想定し、ケース内では太陽指向面を固定している。\n"
             "通信中に構体の太陽面が入れ替わる時系列は解析していない",
             1.60, 3.70, 11.60, 4.20, accent=BLUE, fill=PALE_BLUE,
             title_size=27, body_size=23, centered=False, body_offset=1.25)
    add_card(slide, "ジンバルの場合",
             "構体姿勢がほぼ一定なら、このケース行列の中では同じ扱いになる。\n"
             "熱はバス側の温度場で決まるため",
             14.00, 3.70, 11.40, 4.20, accent=GREEN, fill=PALE_GREEN,
             title_size=27, body_size=23, centered=False, body_offset=1.25)
    add_card(slide, "未解析",
             "相手を追い続ける body pointing の時系列、および軌道の途中で発熱機器の"
             "電源が入る過渡。どちらも温度場の履歴が変わるため、別に解析が要る",
             1.60, 8.40, 23.80, 3.20, accent=ORANGE, fill=PALE_ORANGE,
             title_size=27, body_size=24, centered=False, body_offset=1.25)
    set_notes(slide, "姿勢の頻度を聞かれたときの枚。ケース内で太陽面を固定している前提を"
                     "先に言い、変わる場合は別解析が要ると答える。")
    return slide


def build_b11(prs):
    slide = new_slide(prs, "補足", "B11｜残差Fourier更新の数値（予備）",
                      title_size=32)
    add_table(slide, [
        ["条件", "全区間", "初周", "周1以降"],
        ["Case 13（MY）全誤差を直接Fourier", "1.39 s", "3.37 s", "0.39 s"],
        ["Case 13（MY）階層モデル＋残差Fourier", "0.79 s", "1.60 s", "0.38 s"],
        ["Case 16（PY）全誤差を直接Fourier", "13.2 s", "39.5 s", "0.38 s"],
        ["Case 16（PY）階層モデル＋残差Fourier", "0.76 s", "1.47 s", "0.40 s"],
    ], 1.60, 3.60, 23.40, 5.00, col_widths=[2.4, 0.8, 0.8, 0.9], size=23,
        header_size=23)
    add_text(slide, "Case 16 の直接Fourier は成功率 98.7%。他は 100%",
             1.60, 8.85, 23.40, 0.50, size=21, color=GRAY)
    add_card(slide, "定数の残差更新との差",
             "Case 13 の周1以降は 1.66 → 1.35 s までしか落ちない。軌道周期の山が残る",
             1.60, 9.80, 11.60, 2.70, accent=ORANGE, fill=PALE_ORANGE,
             title_size=25, body_size=22, centered=False, body_offset=1.15)
    add_card(slide, "限界",
             "2ケースのみ。causal、K=2。60秒刻みの全点を使っており、"
             "実センサの疎な成功点では未検証",
             14.00, 9.80, 11.40, 2.70, accent=RED, fill=PALE_RED,
             title_size=25, body_size=22, centered=False, body_offset=1.15)
    add_text(slide, "熱モデルの係数は更新していない。捕捉後の全指向残差に対する予備実験",
             1.60, 12.80, 23.40, 0.60, size=23, color=GRAY,
             align=PP_ALIGN.CENTER)
    set_notes(slide, "5月に構想した適応補正がどこまで進んだかを聞かれたときの枚。"
                     "本編Slide 21の数値の全文。")
    return slide


def build_b12(prs):
    slide = new_slide(prs, "補足", "B12｜前に聞いた秒と、今日の秒の対応",
                      title_size=32)
    add_table(slide, [
        ["いつ", "何の秒だったか", "今日の秒と違う理由"],
        ["2026-05 全体輪講", "仮想の熱バイアスでの捕捉時間と成功率",
         "熱構造解析ではなく、置いた値での試算だった"],
        ["2026-07 光通信RG", "密な走査での捕捉時間",
         "点間隔と検出半径が通信ビーム級で、走査の格子に穴があった"],
        ["2026-08-16 中間メモ", "17ケース集計での捕捉時間",
         "提出稿はCOLD標準表面の14ケースに絞ったあとの数字が正"],
        ["今日（提出稿）", "14.6 → 0.10 s、19.2 → 5.45 s",
         "±1600 µrad・120 µradステップ・検出半径150 µrad、14ケース平均"],
    ], 1.60, 3.60, 23.40, 6.60, col_widths=[1.0, 1.5, 2.6], size=22,
        header_size=23, highlight_last=True, body_align=PP_ALIGN.LEFT)
    add_text(slide,
             ["走査条件が違えば秒は変わる。比べてよいのは、同じ走査条件の中での補正方式どうし",
              "旧い秒はスクリーンに出さない。対応は口頭で答える"],
             1.60, 10.80, 23.40, 1.50, size=25, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.25)
    set_notes(slide, "旧い秒を聞かれたときの手元用。5月は2.43 s → 0.74 s・成功率97.5%で"
                     "仮想バイアス。7月光RGは124.6 s と 156.9 → 59.6 s で、点間隔40 µrad・"
                     "検出半径25 µradの密走査。8月16日のメモは12.1 s と 16.3 → 4.75 s で"
                     "17ケース集計。いずれも走査条件かケース集合が今日と違う。")
    return slide


# --------------------------------------------------------------------------
def renumber(prs, labels) -> None:
    for slide, label in zip(prs.slides, labels):
        shape = page_number_shape(slide)
        if shape is None:
            continue  # title/section slides use the master's slide-number field
        retext(shape, label, size=18, color=GRAY, bold=False,
               align=PP_ALIGN.RIGHT)


def build() -> Path:
    prs = Presentation(TEMPLATE)
    prs.core_properties.title = "衛星光通信の粗捕捉に向けた時変熱バイアスのフィードフォワード補正"
    prs.core_properties.subject = "2026-09-28 全体輪講"
    prs.core_properties.author = "高本 英熙"
    prs.core_properties.comments = (
        "260927_labseminar_slide_outline.md と papers/icso/main.typ から作成。"
        "土台は 20260721_optcommrg_takamoto_v3_issl.pptx。"
    )

    main = [
        build_title(prs),
        build_since_may(prs),
        build_problem(prs),
        build_error_budget(prs),
        build_scope(prs),
        build_satellite(prs),
        build_thermal_desktop(prs),
        build_femap(prs),
        build_los(prs),
        build_cases(prs),
        build_observation(prs),
        build_model_requirements(prs),
        build_hierarchical_model(prs),
        build_identification(prs),
        build_accuracy(prs),
        build_case08(prs),
        build_scan_definition(prs),
        build_nonthermal(prs),
        build_thermal_only(prs),
        build_with_nonthermal(prs),
        build_residual_update(prs),
        build_limitations(prs),
        build_summary(prs),
    ]
    backup = [
        build_backup_divider(prs),
        build_b1(prs),
        build_b2(prs),
        build_b3(prs),
        build_b4(prs),
        build_b5(prs),
        build_b6(prs),
        build_b6b(prs),
        build_b7(prs),
        build_b8(prs),
        build_b9(prs),
        build_b10(prs),
        build_b11(prs),
        build_b12(prs),
    ]

    order = main + backup
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    lookup = {id(slide): sid for slide, sid in zip(list(prs.slides), ids)}
    keep = [lookup[id(slide)] for slide in order]
    for sid in ids:
        sldIdLst.remove(sid)
    for sid in keep:
        sldIdLst.append(sid)
    for sid in ids:
        if sid not in keep:
            prs.part.drop_rel(sid.rId)

    labels = [str(i) for i in range(1, len(main) + 1)]
    labels += ["", "B1", "B2", "B3", "B4", "B5", "B6", "B6", "B7", "B8", "B9",
               "B10", "B11", "B12"]
    renumber(prs, labels)

    # the backups are jumped to by number during Q&A, not shown in the run-through
    for slide in backup:
        slide._element.set("show", "0")

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())

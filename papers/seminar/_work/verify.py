"""Check the produced deck against the outline's forbidden numbers and language rules."""
from __future__ import annotations

import io
import re
from pathlib import Path

from pptx import Presentation

DECK = Path(__file__).resolve().parents[1] / "20260928_labseminar_takamoto.pptx"

FORBIDDEN = [
    "2.43", "0.74 s", "93.9", "97.5", "124.6", "4.78", "156.9", "59.6",
    "94.0", "97.1", "12.1 s", "16.3", "4.75", "数十〜数百 µrad", "40 µrad",
    "25 µrad", "°C", "℃",
]
ALLOW_DEGREE = {"基準温度"}  # Femap reference temperature legitimately uses ℃

MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
MATH_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"


def slide_text(slide):
    chunks = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            chunks.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    chunks.append(cell.text)
    for alt in slide.shapes._spTree.findall(MC):
        chunks.append("".join(t.text or "" for t in alt.iter(MATH_T)))
    return "\n".join(chunks)


def main():
    prs = Presentation(str(DECK))
    out = io.StringIO()
    out.write(f"slides: {len(prs.slides)}\n")
    for index, slide in enumerate(prs.slides, start=1):
        text = slide_text(slide)
        hits = [token for token in FORBIDDEN if token in text]
        hits = [h for h in hits
                if not (h in {"°C", "℃"} and any(k in text for k in ALLOW_DEGREE))]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        ascii_words = sorted({w for w in re.findall(r"[A-Za-z][A-Za-z\-_']{3,}", text)})
        out.write(f"\n--- slide {index} | notes={len(notes)} chars\n")
        if hits:
            out.write(f"  FORBIDDEN: {hits}\n")
        if ascii_words:
            out.write(f"  latin: {' '.join(ascii_words)}\n")
    Path(__file__).with_name("verify.txt").write_text(out.getvalue(), encoding="utf-8")
    print("ok")


if __name__ == "__main__":
    main()

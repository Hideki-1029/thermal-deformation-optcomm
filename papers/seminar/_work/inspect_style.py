import io
from pptx import Presentation
out = io.open("papers/seminar/_work/july_style.txt", "w", encoding="utf-8")
prs = Presentation("papers/seminar/20260721_optcommrg_takamoto_v3_issl.pptx")

out.write("### LAYOUT SHAPES\n")
for i, lay in enumerate(prs.slide_layouts):
    out.write("\n[%d] %s\n" % (i, lay.name))
    for sh in lay.shapes:
        t = sh.text_frame.text.strip().replace("\n"," | ")[:60] if sh.has_text_frame else ""
        out.write("   <%s> %r (%.2f,%.2f,%.2f,%.2f) %s\n" % (sh.shape_type, sh.name,
            (sh.left or 0)/914400,(sh.top or 0)/914400,(sh.width or 0)/914400,(sh.height or 0)/914400, t))

out.write("\n\n### RUN FORMATTING (selected slides)\n")
for idx in [1,4,5,8,9,11,12,13,15,17,19,21,22,25,29,30,31,35,36,44]:
    s = prs.slides[idx-1]
    out.write("\n=== Slide %d (layout %s)\n" % (idx, s.slide_layout.name))
    for sh in s.shapes:
        if not sh.has_text_frame: 
            if getattr(sh,'has_table',False) and sh.has_table:
                tb = sh.table
                c = tb.cell(0,0)
                p = c.text_frame.paragraphs[0]
                r = p.runs[0] if p.runs else None
                out.write("  TABLE %r hdr_font=%s sz=%s color=%s fill?\n" % (sh.name, r.font.name if r else None, r.font.size.pt if r and r.font.size else None, (r.font.color.rgb if r and r.font.color and r.font.color.type is not None else None)))
                # cell fills
                from pptx.oxml.ns import qn
                for ri,row in enumerate(tb.rows):
                    fills=[]
                    for ci,cell in enumerate(row.cells):
                        tcPr = cell._tc.find(qn('a:tcPr'))
                        f = 'none'
                        if tcPr is not None:
                            sf = tcPr.find(qn('a:solidFill'))
                            if sf is not None:
                                sc = sf.find(qn('a:srgbClr'))
                                sm = sf.find(qn('a:schemeClr'))
                                f = sc.get('val') if sc is not None else ('scheme:'+sm.get('val') if sm is not None else '?')
                        fills.append(f)
                    out.write("     row%d fills=%s\n" % (ri, fills))
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                col = None
                try:
                    if r.font.color and r.font.color.type is not None:
                        col = str(r.font.color.rgb)
                except Exception:
                    col = 'theme'
                out.write("  %r: %r font=%s sz=%s bold=%s color=%s\n" % (sh.name, r.text[:40], r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, col))
        # shape fill/line
        try:
            from pptx.oxml.ns import qn
            spPr = sh._element.spPr
            sf = spPr.find(qn('a:solidFill'))
            ln = spPr.find(qn('a:ln'))
            fillc = None
            if sf is not None:
                e = sf.find(qn('a:srgbClr'))
                fillc = e.get('val') if e is not None else 'scheme'
            linec = None
            if ln is not None:
                lsf = ln.find(qn('a:solidFill'))
                if lsf is not None:
                    e = lsf.find(qn('a:srgbClr'))
                    linec = e.get('val') if e is not None else 'scheme'
            if fillc or linec:
                out.write("     SHAPE %r fill=%s line=%s w=%s\n" % (sh.name, fillc, linec, ln.get('w') if ln is not None else None))
        except Exception as e:
            pass
out.close()
print("ok")

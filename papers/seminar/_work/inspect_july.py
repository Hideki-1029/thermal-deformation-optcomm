import io, sys
from pptx import Presentation

out = io.open("papers/seminar/_work/july_dump.txt", "w", encoding="utf-8")
prs = Presentation("papers/seminar/20260721_optcommrg_takamoto_v3_issl.pptx")
out.write("SLIDE SIZE: %.2f x %.2f in\n" % (prs.slide_width/914400, prs.slide_height/914400))
out.write("LAYOUTS:\n")
for i, lay in enumerate(prs.slide_layouts):
    out.write("  [%d] %s\n" % (i, lay.name))
out.write("N SLIDES: %d\n" % len(prs.slides))
for i, s in enumerate(prs.slides, 1):
    out.write("\n=== Slide %d | layout=%s | shapes=%d\n" % (i, s.slide_layout.name, len(s.shapes)))
    for sh in s.shapes:
        kind = sh.shape_type
        info = "  <%s> name=%r pos=(%.2f,%.2f) size=(%.2f,%.2f)" % (
            kind, sh.name,
            (sh.left or 0)/914400, (sh.top or 0)/914400,
            (sh.width or 0)/914400, (sh.height or 0)/914400)
        out.write(info + "\n")
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                out.write("      TEXT: " + t.replace("\n", " | ")[:300] + "\n")
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                out.write("      ROW: " + " | ".join(c.text.replace("\n"," ")[:40] for c in r.cells) + "\n")
    if s.has_notes_slide:
        nt = s.notes_slide.notes_text_frame.text.strip()
        if nt:
            out.write("  NOTES: " + nt.replace("\n", " | ")[:400] + "\n")
out.close()
print("done")
